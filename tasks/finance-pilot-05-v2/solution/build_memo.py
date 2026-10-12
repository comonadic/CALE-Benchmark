#!/usr/bin/env python3
"""Create the editable reference memo with declared python-docx only."""
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def build(path):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.top_margin=Cm(2.1);s.bottom_margin=Cm(2.0);s.left_margin=Cm(2.2);s.right_margin=Cm(2.2)
    for name,size in [('Normal',10.5),('Title',20),('Heading 1',14),('Heading 2',11.5)]:
        st=d.styles[name];st.font.name='Noto Sans CJK SC';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.10
        if name=='Heading 1':st.paragraph_format.space_before=Pt(9)
        if name=='Heading 2':st.paragraph_format.space_before=Pt(7)
        for border in st._element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for border in d.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    def p(text):return d.add_paragraph(text)
    def h(text):return d.add_heading(text,1)
    def table(headers,rows,widths):
        t=d.add_table(rows=1,cols=len(headers));t.autofit=False
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for i,v in enumerate(headers):t.rows[0].cells[i].text=v
        for row in rows:
            for cell,text in zip(t.add_row().cells,row):cell.text=text
        pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
        pr.append(b)
        for cell in t.rows[0].cells:
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(fill)
            for run in cell.paragraphs[0].runs:run.bold=True
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4)
                    for run in para.runs:run.font.size=Pt(9.5)
        d.add_paragraph().paragraph_format.space_after=Pt(0)
        return t
    d.add_paragraph('中央预算审计整改事项审阅意见','Title')
    p('致项目负责人    报告信息时点 2025年12月22日')
    p('公开报告已披露资金盘活、归还原渠道和工程推进等进展。建议下一轮优先核对国债资金项目明细与银行、预算账务记录，并分别取得海河和鄱阳湖工程阶段清单及验收资料。目前可形成披露比较和报告勾稽；资金实际余额、资金消化及工程最终闭环仍需底层证据支持。附台账保留11个事项的原值、计算、关系和补证安排。')
    h('总体口径与资金比较')
    p('主报告署日期为2025年12月22日；总览截至2025年9月底，三类问题合计2186+1299+753=4238个，98%仅适用于2186个立行立改问题。结语另披露62个（占2%）尚未整改到位，未单列该计数截止日。两处原值均保留。[F048物理1、3、22/印刷1、3、22 一（三）及结语首段]')
    p('单从整数披露的98%不能反推精确完成数。假设62与2186同截止、同分类和同分母，则62/2186=2.8362%，条件完成2186−62=2124个、2124/2186=97.1638%；通常整数百分比舍入为3%和97%，不能与原文2%和98%勾稽。现有资料不能确认时点、分类或分母差异，也不能认定文字差错；台账的已确认精确完成数留空，2124仅作条件复算。共整改问题金额1.04万亿元为10400亿元，按问题整改汇总口径保留。')
    table(['事项','原问题','已报告进展','披露比较'],[
      ['N01 超需求闲置','267亿元\n10省115项目','34.65亿元\n6省盘活使用','差额232.35亿元\n覆盖12.98%'],
      ['N01a 沉淀子项','48亿元\n3省31项目','未单列','属于267亿元的其中项'],
      ['N02 挤占挪用','57亿元\n13省','44.66亿元\n13省整改','差额12.34亿元\n覆盖78.35%'],
      ['N03 违规招标等','12省120项目','10省整改涉及资金\n68.85亿元','缺原问题金额与整改项目数'],
    ],[3.7,3.6,4.0,5.3])
    p('N01覆盖比为34.65/267×100，原问题10省与盘活6省需项目匹配；N02为44.66/57×100，即使均为13省仍需逐笔同范围核对。232.35亿元、12.34亿元是披露差额，不是已验证的损失或实际未归还余额。48亿元已在267亿元内，且不能直接把34.65亿元全部分配到31个子项。68.85亿元没有同口径原金额分母，也不能以120项目作金额分母。[F049物理3–4/印刷2–3 二（二）；F048物理8–9/印刷8–9 二（三）2]')
    section_title=h('工程阶段与部门项目');section_title.paragraph_format.page_break_before=True
    p('海河流域有关7省的264个原未完成任务河段，披露已开工78、完工8、前期178，按列示三个阶段相加为264。阶段占比分别为29.55%、3.03%、67.42%。开工78个和完工8个在本次三阶段核对中分别记录；公开完工8个也需要竣工验收与唯一河段编号支持核查。[F049物理3/印刷2 二（一）；F048物理8/印刷8 二（三）1]')
    p('鄱阳湖原134个项目计划总投资120.3亿元、实际完成投资7.5%是原问题时点资料。51个逾期项目中43个已开（完）工，合并阶段覆盖43/51×100=84.31%，其余8个未列入该类别。43项没有披露开工与完工的拆分，因此不能视为43项全部竣工；历史投资比例也不能作为新的整改投资完成率。[F049物理4–5/印刷3–4 二（三）；F048物理17/印刷17 生态环境保护修复第二方面]')
    p('附件1在开篇将有关闲置及挪用问题定位为2023年增发国债资金。水利部湿地蓄滞洪区项目为1个项目的1.86亿元超长期特别国债闲置，优化布局后报告2025年7月全面开工。它与全国增发国债事项只有同次发布及主题关联；全面开工支持开工动作已报告完成，未提供1.86亿元实际消化金额及资金余额。[F049物理2/印刷1 二；F050物理14/印刷14 序号9首行]')
    p('水利知识平台由中国水利水电出版传媒集团有限公司在2023至2024年申请国有资本经营预算5000万元，项目获批后进展缓慢，闲置4420.05万元。两者分别为0.5亿元、0.442005亿元。建立工作台账、加快消化和已申请调减2025年预算属于所披露整改措施；申请不等于调减已获批准，也没有已消化资金数字。预算与闲置金额之差不能直接充作已支出或整改金额。[F050物理15/印刷15 序号9水利知识平台行]')
    h('汇总关系与舍入差异')
    p('主报告灾后恢复重建和防灾减灾能力提升节汇总为148.17亿元，三个披露分项34.65+44.66+68.85=148.16亿元，差0.01亿元，即100万元。原汇总和分项均保留。四个两位小数披露值的舍入可形成这种差异，保守误差界为0.02亿元；未舍入数据、汇总调整或未列示项目也需核实，现有资料不能唯一确定成因或确认漏整改100万元。此项仅为报告层勾稽，主附件重述不重复相加，部门两项目及不同问题金额也不建立新的总风险敞口。[F048物理8–9/印刷8–9 二（三）2；F049物理4/印刷3 二（二）]')
    section_title=h('优先补证与闭环判据');section_title.paragraph_format.page_break_before=True
    p('以下期限均为本次审阅建议，按责任部门收到清单后起算。先完成项目与资金唯一编号匹配，再确定实际余额及闭环；涉及工程的项目同时核对资金和实物工作量。')
    actions=[
      ('N01与N01a 国债闲置及沉淀','国家发展改革委、财政部督促有关地方，建议10个工作日内提供115项目资金到账、支付、盘活与进度明细；3省补31项目预付款、施工单位银行余额、退款或抵扣与实物工作量。按原项目编号核对48在267内及34.65覆盖哪些项目，凭到账、支付与进度相互印证判定资金消化。'),
      ('N02 挤占挪用','13省财政部门建议10个工作日内提供逐笔挪用、归还原渠道的银行回单、预算账务调整及余额对账。以57亿元原清单匹配44.66亿元整改记录，核对金额、渠道、日期和用途，核定真实待处理余额。'),
      ('N03 违规招标等','有关地方招投标监管部门建议15个工作日内提供120项目问题与处理清单、重新招标或合同纠正资料、责任追究记录及资金口径。以项目处理决定和整改实施凭证确认程序问题已纠正，并建立68.85亿元与原项目的对应。'),
      ('N04 海河河段','水利部及7省水利主管部门建议15个工作日内提交264河段唯一编号和分阶段清单、78项开工记录、8项完工验收、178项前期要件与责任节点。核对阶段互斥、验收范围和后续节点，逐段判断实际完成。'),
      ('N05 鄱阳湖项目','项目地方主管部门建议15个工作日内提交51项目分别开工、完工及验收清单、最新已完成投资和8项目状态。依据项目编号核对43合并数，凭竣工验收判定完工，投资对比使用一致期间和计划范围。'),
      ('D01与D02 部门项目','水利部及湿地项目地方单位建议10个工作日内补调整后环评批复、开工记录、支付与余额；出版传媒集团建议15个工作日内补招采支付台账、银行余额、可行性复核和2025预算调减申请及批复。分别验证合规开工、实际消化及批准收回，不将申请和施工启动视为全资金闭环。'),
      ('O01 O02 R01 汇总核对','审计整改牵头方与报告汇总部门建议10个工作日内补同期分类问题台账、逐项销号及统计口径，说明第3页98%与第22页62个（占2%）的分母、截止与分类；另补未舍入金额、互斥项目标识和汇总调整表，以底表核对差异和重复覆盖。'),
    ]
    for title,text in actions:
        d.add_heading(title,2);p(text)
    d.core_properties.title='中央预算审计整改事项审阅意见';d.core_properties.author='财政资金审阅团队'
    d.core_properties.subject='资金与工程证据追踪及补证安排'
    Path(path).parent.mkdir(parents=True,exist_ok=True);d.save(path)

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);a=p.parse_args();build(a.output)
