"""Native-office numeric/dependency upper bounds, never a final semantic score.

Locations use the public model-data-dictionary named-region handoff contract.
No reference sheet names, row numbers or exact formula spellings are required.
"""
from __future__ import annotations
import argparse, copy, json, math, os, re, shutil, subprocess, tempfile, time, traceback
from pathlib import Path
import openpyxl
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries
from docx import Document
from pypdf import PdfReader

ANCHOR=92058481.99; REVENUE=79972916.47; SHARES=800000000
TOL=.05
AUTO_WEIGHTS={'R01':14,'R02':20,'R03':20,'R04':13,'R05':10,'R07':3}
SUB_WEIGHTS={'R01a':14,'R02a':10,'R02b':6,'R02c':4,'R03a':7,'R03b':7,'R03c':6,'R04a':5,'R04b':4,'R04c':4,'R05a':5,'R05b':5,'R07a':3}
def subcriterion(e):
    g=e['criterion'];i=e['id']
    if g=='R01':return 'R01a'
    if g=='R02':return 'R02c' if 'variance' in i else 'R02b' if 'project_cash' in i or 'scope_difference' in i else 'R02a'
    if g=='R03':
        if i.startswith('invalid-') or 'unselected-blanks' in i or 'proposed-zero' in i or 'exact-D-zero' in i:return 'R03c'
        if i.startswith('dependencies-') or i.startswith('editable-') or 'dependency-perturb-' in i or 'historical-invariant' in i or i.startswith('native-formula-'):return 'R03b'
        return 'R03a'
    if g=='R04':return 'R04c' if 'budget' in i or 'remaining' in i else 'R04b' if 'margin' in i else 'R04a'
    if g=='R05':return 'R05a' if i.startswith('initial-') or i.startswith('sensitivity-count-') else 'R05b'
    return 'R07a'
# Expected facts; positions and labels are solely the public named-region contract.
SPECS={'np': 23280783.76,
 'da': 41849395.41,
 'interest': 3427256.05,
 'tax': 4242395.02,
 'ebitda_control': 72799830.24,
 'wc': 60480597,
 'other': 2390988185.77,
 'capex': -2186382254.09,
 'principal': -200000000,
 'cash_interest_tax': -7104333.27,
 'old_reserve': -38723543.66,
 'd_control': 92058481.99,
 'revenue': 79972916.47,
 'project_ebitda': 68120799.93,
 'cash_in': 213952232.49,
 'cash_out': 95638012.92,
 'cash_control': 118314219.57,
 'fund_cash': 126973070.63,
 'fund_income': 89457272.23,
 'shares': 800000000,
 'actual_paid': 0,
 'ratio': 0.9,
 'management_fee': 2196163.56,
 'custody_fee': 109807.51,
 'operator_fee': 9964029.05}
OUTPUTS=['selector', 'custom_decline', 'custom_reserve', 'proposed_input', 'D', 'lower', 'upper', 'positive', 'per_unit', 'proposed', 'upper_margin', 'lower_margin', 'decision', 'budget', 'remaining', 'budget_gap', 'remaining_gap', 'historical', 'ebitda', 'pre_reserve', 'project_cash', 'scope_difference', 'historical_per_unit', 'ebitda_variance', 'd_variance', 'cash_variance']
def numeric(v): return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)
def near(v,e,t=TOL):return numeric(v) and abs(v-e)<=t
def unavailable(v):return v is None or (isinstance(v,str) and not v.startswith('='))
class CandidateContractError(ValueError):
    """A candidate's public named-region handoff is missing or malformed."""

class Evaluation:
    def __init__(self, workbook, logdir, office):
        self.path=Path(workbook);self.logdir=Path(logdir);self.office=office
        self.events=[];self.mapping={};self.scenarios=[]
        self.wb=openpyxl.load_workbook(self.path, data_only=False)
        self.contract=json.loads(Path(__file__).with_name('model-data-dictionary.json').read_text())
        fields=self.contract['scalar_fields'];points=self.contract['sensitivity_points']
        if self.contract['version']!=1 or {x['key'] for x in fields}!=set(SPECS)|set(OUTPUTS) or len(points)!=30:
            raise RuntimeError('Trusted business data dictionary disagrees with verifier schema')
        self.locations={}
        for field in fields:self.locations[field['key']]=self.find(field['name'],optional=not field['required'])
        for point in points:self.locations[point['name']]=self.find(point['name'])
        self.global_scale=self.scale('metric_amount_scale')
        self.share_scale=self.scale('metric_share_scale')
        self.cached=self.recalculate(self.wb,'initial')
        self.locate()
        self.sensitivity=self.locate_sensitivity()
    def record(self,criterion,id,actual,expected,passed=None,tolerance=None,details=None):
        if passed is None: passed=near(actual,expected,tolerance or TOL) if numeric(expected) else actual==expected
        self.events.append(dict(criterion=criterion,id=id,passed=bool(passed),actual=actual,expected=expected,tolerance=tolerance,details=details))
    def recalculate(self,wb,tag):
        configroot=Path(os.environ.get('OFFICE_CONFIG_ROOT','/tmp/harbor-office'));configroot.mkdir(parents=True,exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='finance03-',dir=configroot) as tmp:
            tmp=Path(tmp);source=tmp/'source';dest=tmp/'recalculated';source.mkdir();dest.mkdir()
            file=source/'model.xlsx';wb.save(file)
            profile=tmp/'profile'
            cmd=[self.office,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','xlsx','--outdir',str(dest),str(file)]
            p=subprocess.run(cmd,capture_output=True,text=True,timeout=45)
            out=dest/'model.xlsx'
            self.scenarios.append({'tag':tag,'exit_code':p.returncode,'stdout':p.stdout[-1500:],'stderr':p.stderr[-1500:]})
            if p.returncode or not out.exists():raise RuntimeError('Office recalculation failed: '+p.stdout+p.stderr)
            return openpyxl.load_workbook(out,data_only=True)
    def find(self,name,optional=False):
        matches=[x for key,x in self.wb.defined_names.items() if key.lower()==name.lower()]
        locals_found=[x for sh in self.wb for key,x in sh.defined_names.items() if key.lower()==name.lower()]
        if not matches and optional and not locals_found:return None
        if len(matches)!=1 or locals_found:raise CandidateContractError(f'{name}: one workbook-scope named region required; missing, duplicate or local-only name')
        defined=matches[0]
        try:
            if defined.type!='RANGE':raise ValueError('direct region required')
            destinations=list(defined.destinations)
            if len(destinations)!=1:raise ValueError('one destination required')
            sheet,address=destinations[0]
            sheet=next((s for s in self.wb.sheetnames if s.lower()==sheet.lower()),sheet)
            if sheet not in self.wb.sheetnames:raise ValueError('sheet does not exist')
            minc,minr,maxc,maxr=range_boundaries(address.replace('$','').upper())
            if None in [minc,minr,maxc,maxr] or minc!=maxc or minr!=maxr or not 1<=minc<=16384 or not 1<=minr<=1048576:raise ValueError('one XLSX cell required')
            return sheet,self.wb[sheet].cell(minr,minc).coordinate
        except (TypeError,ValueError,KeyError,IndexError) as exc:raise CandidateContractError(f'{name}: invalid named region ({exc})') from exc
    def scale(self,name):
        loc=self.find(name,optional=True)
        if loc is None:return 1
        value=self.wb[loc[0]][loc[1]].value
        if not numeric(value) or value<=0:raise CandidateContractError(f'{name}: finite positive typed unit multiplier required')
        return value
    def locate(self):
        for k,expected in SPECS.items():
            loc=self.locations[k]
            if loc is None:continue # narrative-only policy/fee evidence receives professional review
            raw=self.cached[loc[0]][loc[1]].value
            factor=1 if k=='ratio' else self.share_scale if k=='shares' else self.global_scale
            self.mapping[k]={'location':loc,'factor':factor,'type':'source','raw':raw}
            observed=raw*factor if numeric(raw) else raw
            valid=near(observed,expected,1e-8 if k=='ratio' else TOL)
            if expected<0:valid=near(abs(observed),abs(expected)) # positive deductions with subtraction are legal
            self.record('R01','source-'+k,observed,expected,valid,1e-8 if k=='ratio' else TOL)
        for k in OUTPUTS:
            kind='input' if k in ['selector','custom_decline','custom_reserve','proposed_input'] else 'formula'
            loc=self.locations[k]
            if loc is None and k=='positive':loc=self.mapping.get('upper',{}).get('location')
            if loc is None and k=='proposed':loc=self.mapping.get('proposed_input',{}).get('location')
            if loc:
                factor=1 if k in ['selector','custom_decline','decision','per_unit','historical_per_unit'] else self.global_scale
                self.mapping[k]={'location':loc,'factor':factor,'type':kind}
                if k=='selector' and self.wb[loc[0]][loc[1]].value in ['基准','压力','自定义']:
                    self.mapping[k]['case_style']='names'
                if kind=='input':self.record('R03','editable-'+k,self.wb[loc[0]][loc[1]].data_type!='f',True)
                elif k not in ['positive','proposed']:
                    group='R02' if k in ['historical','ebitda','pre_reserve','project_cash','scope_difference','historical_per_unit','ebitda_variance','d_variance','cash_variance'] else 'R03' if k=='D' else 'R04'
                    self.record(group,'native-formula-'+k,self.wb[loc[0]][loc[1]].data_type=='f',True)
        # Optional independent historical per-unit can be assessed from historical D and shares.
    def get(self,wb,key):
        m=self.mapping[key];v=wb[m['location'][0]][m['location'][1]].value
        return v*m['factor'] if numeric(v) else v
    def set(self,wb,key,value):
        m=self.mapping[key]
        if key=='selector' and m.get('case_style')=='names' and value in [1,2,3]:value=['基准','压力','自定义'][int(value)-1]
        wb[m['location'][0]][m['location'][1]]=value/m['factor'] if numeric(value) else value
    def deps(self,loc,seen=None):
        seen=set() if seen is None else seen
        if loc in seen:return seen
        seen.add(loc);val=self.wb[loc[0]][loc[1]].value
        if not isinstance(val,str) or not val.startswith('='):return seen
        try:tokens=Tokenizer(val).items
        except Exception:return seen
        for token in tokens:
            if token.type!='OPERAND' or token.subtype!='RANGE':continue
            ref=token.value
            local=self.wb[loc[0]].defined_names
            name=next((v for k,v in local.items() if k.lower()==ref.lower()),None) or next((v for k,v in self.wb.defined_names.items() if k.lower()==ref.lower()),None)
            refs=list(name.destinations) if name else [(None,ref)]
            for sheet,addr in refs:
                if '!' in addr:sheet,addr=addr.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
                sheet=sheet or loc[0]
                sheet=next((s for s in self.wb.sheetnames if s.lower()==sheet.lower()),sheet)
                if sheet not in self.wb.sheetnames:continue
                try:
                    minc,minr,maxc,maxr=range_boundaries(addr.replace('$','').upper())
                    if None in [minc,minr,maxc,maxr] or (maxc-minc+1)*(maxr-minr+1)>10000:continue
                    for row in self.wb[sheet].iter_rows(min_row=minr,max_row=maxr,min_col=minc,max_col=maxc):
                        for cell in row:self.deps((sheet,cell.coordinate),seen)
                except ValueError:continue
        return seen
    def dependency(self,key,required,criterion):
        loc=self.mapping[key]['location'];deps=self.deps(loc)
        missing=[x for x in required if not any(self.mapping[v]['location'] in deps for v in x.split('|') if v in self.mapping)]
        unsupported=any('INDIRECT(' in str(self.wb[s][c].value).upper() or 'OFFSET(' in str(self.wb[s][c].value).upper() for s,c in deps)
        if missing and unsupported:
            # Native upstream perturbations remain scored; the bounded static parser does not reject volatile equivalent formulas.
            self.scenarios.append({'tag':'static-trace-limited-'+key,'unsupported':'INDIRECT/OFFSET','missing_static_metrics':missing,'fallback':'native dynamic upstream perturbations'})
            return
        self.record(criterion,'dependencies-'+key,missing,[],details={'cell':loc,'dependency_count':len(deps),'required_metrics':required})
    def locate_sensitivity(self):
        found={'D':{},'margin':{}}
        for point in self.contract['sensitivity_points']:
            key=f"{float(point['revenue_decline'])}:{float(point['extra_reserve_yuan'])}"
            found[point['kind']][key]=self.locations[point['name']]
        return found
    def sensitivity_checks(self,wb,proposed,tag,anchor=ANCHOR,revenue=REVENUE):
        for a in [0,.05,.1,.15,.2]:
            for b in [0,5e6,1e7]:
                key=f'{float(a)}:{float(b)}'
                # Keys stringify consistently with the rounded float values used above.
                for kind in ['D','margin']:
                    loc=self.sensitivity[kind].get(key)
                    e=anchor-revenue*a-b if kind=='D' else max(anchor-revenue*a-b,0)-proposed if numeric(proposed) else 'unavailable'
                    val=wb[loc[0]][loc[1]].value if loc else 'missing'
                    val=val*self.global_scale if numeric(val) else val
                    self.record('R05',f'{tag}-sensitivity-{kind}-{a}-{b}',val,e,unavailable(val) and loc is not None if e=='unavailable' else None)
    def check_budget(self,wb,key,signed,tag):
        v=self.get(wb,key)
        passed=near(v,signed)
        details={'accepted':'signed budget or zero floor with linked explicit shortfall'}
        if not passed and signed<0 and near(v,0):
            gapkey='budget_gap' if key=='budget' else 'remaining_gap'
            if gapkey in self.mapping:
                gap=self.get(wb,gapkey);passed=near(gap,-signed)
                self.dependency(gapkey,['proposed_input','revenue'],'R04')
                details['shortfall']=gap
        self.record('R04',tag+'-'+key,v,signed,passed,details=details)
    def active(self,wb,tag,decline,reserve,proposed,anchor=ANCHOR,revenue=REVENUE,shares=SHARES):
        D=anchor-revenue*decline-reserve;U=max(D,0);L=.9*U
        expected={'D':D,'upper':U,'lower':L,'positive':U,'per_unit':U/shares,'proposed':proposed,'upper_margin':U-proposed,'lower_margin':proposed-L}
        for k,e in expected.items():self.record('R03' if k=='D' else 'R04',tag+'-'+k,self.get(wb,k),e,tolerance=1e-8 if k=='per_unit' else TOL)
        state='none' if D<=0 else 'below' if proposed<L-TOL else 'above' if proposed>U+TOL else 'within'
        actual=self.get(wb,'decision')
        # A nonempty native status can support only an upper bound. The independent
        # judge must read this actual cell in this recalculated case, including
        # negation, equality and D<=0. No vocabulary list assigns semantic credit.
        self.record('R04',tag+'-decision',actual,state,isinstance(actual,str) and bool(actual.strip()),details={
            'assessment':'objective_upper_bound_only; status meaning pending independent rubric',
            'cell':self.mapping['decision']['location'],'scenario_tag':tag,
            'D':D,'positive_D':U,'lower':L,'upper':U,'proposed':proposed,
            'at_lower_bound':D>0 and abs(proposed-L)<=TOL,
            'at_upper_bound':D>0 and abs(proposed-U)<=TOL})
        self.check_budget(wb,'budget',anchor-revenue*decline-proposed,tag)
        self.check_budget(wb,'remaining',D-proposed,tag)
        self.record('R03',tag+'-historical-invariant',self.get(wb,'historical'),anchor)
    def run(self):
        for k,e in {'ebitda':72799830.24,'pre_reserve':130782025.65,'historical':ANCHOR,'project_cash':118314219.57,'scope_difference':8658851.06,'ebitda_variance':0,'d_variance':0,'cash_variance':0}.items():
            self.record('R02','history-'+k,self.get(self.cached,k),e)
        if 'historical_per_unit' in self.mapping:self.record('R02','history-per-unit',self.get(self.cached,'historical_per_unit'),ANCHOR/SHARES,tolerance=1e-8)
        for key,required,criterion in [
            ('ebitda',['np','da','interest','tax'],'R02'),
            ('historical',['np','da','interest','tax','wc','other','capex','principal','cash_interest_tax','old_reserve'],'R02'),
            ('project_cash',['cash_in','cash_out'],'R02'),
            ('scope_difference',['fund_cash','cash_control|project_cash'],'R02'),
            ('D',['historical','revenue','selector','custom_decline','custom_reserve'],'R03'),
            ('upper_margin',['proposed_input','D'],'R04'),('lower_margin',['proposed_input','D'],'R04'),
            ('per_unit',['shares','D'],'R04'),('budget',['historical','revenue','proposed_input'],'R04')]:self.dependency(key,required,criterion)
        for kind in self.sensitivity:
            self.record('R05','sensitivity-count-'+kind,len(self.sensitivity[kind]),15)
            for key,loc in self.sensitivity[kind].items():
                deps=self.deps(loc);req=['historical','revenue']+(['proposed_input'] if kind=='margin' else [])
                self.record('R05','sensitivity-dependencies-'+kind+'-'+key,all(self.mapping[x]['location'] in deps for x in req),True)
        self.sensitivity_checks(self.cached,85e6,'initial')
        cases=[(1,0,0),(2,.1,5e6),(3,.2,1e7)]
        for selector,decline,reserve in cases:
            wb=openpyxl.load_workbook(self.path);self.set(wb,'selector',selector)
            rec=self.recalculate(wb,f'case-{selector}');self.active(rec,f'case-{selector}',decline,reserve,85e6)
        mutations=[
          ('custom-change',3,.3,30e6,85e6),('negative-D',3,1,100e6,85e6),
          ('below-90',1,.2,1e7,80e6),('proposed-zero',1,.2,1e7,0),
          ('exact-D-zero-proposed-zero',3,0,ANCHOR,0),('exact-D-zero-positive-proposed',3,0,ANCHOR,85e6),
          ('at-lower-bound',1,.2,1e7,.9*ANCHOR),('at-upper-bound',1,.2,1e7,ANCHOR),
        ]
        for tag,selector,decline,reserve,proposed in mutations:
            wb=openpyxl.load_workbook(self.path)
            native_anchor=self.get(self.cached,'historical')
            if tag.startswith('exact-D-zero'):reserve=native_anchor
            if tag=='at-upper-bound':proposed=native_anchor
            if tag=='at-lower-bound':proposed=.9*native_anchor
            for key,x in [('selector',selector),('custom_decline',decline),('custom_reserve',reserve),('proposed_input',proposed)]:self.set(wb,key,x)
            rec=self.recalculate(wb,tag)
            self.active(rec,tag,decline if selector==3 else 0,reserve if selector==3 else 0,proposed,anchor=native_anchor if tag.startswith('exact-D-zero') else ANCHOR)
            if tag in ['below-90','proposed-zero']:self.sensitivity_checks(rec,proposed,tag)
        invalid=[('blank-proposed','proposed_input',None,1),('negative-proposed','proposed_input',-1,1),('text-proposed','proposed_input','pending',1),
                 ('blank-custom-decline','custom_decline',None,3),('decline-too-high','custom_decline',1.1,3),('decline-negative','custom_decline',-.1,3),
                 ('blank-custom-reserve','custom_reserve',None,3),('negative-reserve','custom_reserve',-1,3),('text-reserve','custom_reserve','pending',3),
                 ('blank-selector','selector',None,None),('invalid-selector','selector',4,None),('fractional-selector','selector',1.5,None)]
        for tag,key,x,selector in invalid:
            wb=openpyxl.load_workbook(self.path)
            if selector is not None:self.set(wb,'selector',selector)
            self.set(wb,key,x);rec=self.recalculate(wb,tag)
            relevant=['proposed','upper_margin','lower_margin','budget','remaining'] if key=='proposed_input' else ['D','upper','lower','per_unit']
            for metric in relevant:self.record('R03','invalid-'+tag+'-'+metric,self.get(rec,metric),'unavailable',unavailable(self.get(rec,metric)))
            actual=self.get(rec,'decision')
            self.record('R03','invalid-'+tag+'-status',actual,'input unavailable',isinstance(actual,str) and bool(actual.strip()),details={
                'assessment':'objective_upper_bound_only; status meaning pending independent rubric',
                'cell':self.mapping['decision']['location'],'scenario_tag':tag,
                'mutated_metric':key,'mutated_value':x,'selected_case':selector,
                'required_meaning':'Active input unavailable; explain missing, nonnumeric or out-of-range input without treating it as zero.'})
            if key=='proposed_input':self.sensitivity_checks(rec,None,tag)
        wb=openpyxl.load_workbook(self.path);self.set(wb,'selector',1);self.set(wb,'custom_decline',None);self.set(wb,'custom_reserve',None)
        self.active(self.recalculate(wb,'unselected-blanks'),'unselected-blanks',0,0,85e6)
        # Perturb upstream actuals, not only scenario inputs, to detect copied disclosed answers.
        for key,delta,tag in [('np',1234.56,'profit'),('old_reserve',2345.67,'existing-reserve'),('revenue',7777.77,'project-revenue'),('shares',1e6,'units')]:
            wb=openpyxl.load_workbook(self.path);self.set(wb,'selector',2);original=self.get(self.cached,key);self.set(wb,key,original+delta)
            rec=self.recalculate(wb,'dependency-perturb-'+tag)
            anchor=ANCHOR
            if key=='np':anchor+=delta
            if key=='old_reserve':anchor+=delta if original<0 else -delta
            self.active(rec,'dependency-perturb-'+tag,.1,5e6,85e6,anchor=anchor,revenue=REVENUE+delta if key=='revenue' else REVENUE,shares=SHARES+delta if key=='shares' else SHARES)
            self.sensitivity_checks(rec,85e6,'dependency-perturb-'+tag,anchor=anchor,revenue=REVENUE+delta if key=='revenue' else REVENUE)
        # Formula errors are checked against recalculated native results, not regex over formulas.
        errors=[f'{sh.title}!{c.coordinate}:{c.value}' for sh in self.cached for row in sh for c in row if c.data_type=='e']
        self.record('R07','recalculated-native-formula-errors',errors,[])
        return self.events

def document_checks(path,office,logdir):
    events=[]
    def add(id,actual,expected,passed):events.append(dict(criterion='R07',id=id,actual=actual,expected=expected,passed=passed,tolerance=None,details=None))
    doc=Document(path)
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    add('native-docx-readable-substantive',len(text),'nonempty professional note',len(text)>=300)
    configroot=Path(os.environ.get('OFFICE_CONFIG_ROOT','/tmp/harbor-office'));configroot.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='finance03-doc-',dir=configroot) as tmp:
        tmp=Path(tmp);p=subprocess.run([office,'-env:UserInstallation='+(tmp/'profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(tmp),str(path)],capture_output=True,text=True,timeout=45)
        pdf=tmp/(Path(path).stem+'.pdf')
        if p.returncode or not pdf.exists():raise RuntimeError('DOCX page rendering failed: '+p.stdout+p.stderr)
        pages=len(PdfReader(pdf).pages);add('docx-four-page-limit',pages,'1 to 4 pages',1<=pages<=4)
        shutil.copy(pdf,Path(logdir)/'submitted-note.pdf')
    return events

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--output-dir',default='/workspace/output');ap.add_argument('--log-dir',default='/logs/verifier');ap.add_argument('--office',default=os.environ.get('SOFFICE_BIN','soffice'));ap.add_argument('--workbook-only',action='store_true');args=ap.parse_args()
    logdir=Path(args.log_dir);logdir.mkdir(parents=True,exist_ok=True)
    (logdir/'reward.txt').write_text('0\n')
    import hashlib
    report={'version':4,'verifier_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'executed_at':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'status':'failed','reward_scope':'normalized objective technical upper bounds only; semantic and professional scores are pending independent full rubric review','automatic_points_available':80,'manual_points_pending':20,'professional_score':None,'final_score':None,'status_semantics_assessed':False,'events':[]}
    try:
        out=Path(args.output_dir);e=Evaluation(out/'distribution-model.xlsx',logdir,args.office)
        report['public_contract_sha256']=hashlib.sha256(Path(__file__).with_name('model-data-dictionary.json').read_bytes()).hexdigest()
        report['events']=e.run();report['mapping']=e.mapping;report['sensitivity_mapping']=e.sensitivity;report['native_engine_runs']=e.scenarios
        if not args.workbook_only:report['events']+=document_checks(out/'distribution-note.docx',args.office,logdir)
        groups={};subgroups={}
        for event in report['events']:event['subcriterion']=subcriterion(event)
        for criterion,weight in SUB_WEIGHTS.items():
            checks=[x for x in report['events'] if x['subcriterion']==criterion]
            if args.workbook_only and criterion=='R07a':weight=1
            passed=sum(x['passed'] for x in checks)
            subgroups[criterion]={'automatic_points_available':weight,'passed_checks':passed,'total_checks':len(checks),'automatic_points_earned':weight*passed/len(checks) if checks else 0}
        for criterion,weight in AUTO_WEIGHTS.items():
            parts=[v for k,v in subgroups.items() if k.startswith(criterion)]
            groups[criterion]={key:sum(p[key] for p in parts) for key in ['automatic_points_available','passed_checks','total_checks','automatic_points_earned']}
        total=sum(g['automatic_points_available'] for g in groups.values());earned=sum(g['automatic_points_earned'] for g in groups.values())
        reward=earned/total;report.update(criterion_summary=groups,subcriterion_summary=subgroups,automatic_points_available=total,automatic_points_earned=earned,objective_technical_points_upper_bound=earned,reward=reward,status='passed' if all(x['passed'] for x in report['events']) else 'failed',failed_checks=[x['id'] for x in report['events'] if not x['passed']])
        (logdir/'reward.txt').write_text(f'{reward:.8f}\n')
    except CandidateContractError as exc:
        report.update(status='invalid_submission',candidate_failure_reason=str(exc),reward=0,automatic_points_earned=0,objective_technical_points_upper_bound=0,final_score=0,grading_complete=True,failed_checks=['public-named-region-handoff'])
        report['events']=[{'criterion':'R07','id':'public-named-region-handoff','passed':False,'actual':str(exc),'expected':'complete valid public business named-region index'}]
        report['criterion_summary']={k:{'automatic_points_available':v,'automatic_points_earned':0,'passed_checks':0,'total_checks':1} for k,v in AUTO_WEIGHTS.items()}
        (logdir/'reward.txt').write_text('0\n')
    except Exception as exc:
        report['error']=str(exc);report['traceback']=traceback.format_exc();report['reward']=0
        report['verifier_failed']=True
        report['events'].append({'criterion':'R07','id':'verifier-execution','passed':False,'actual':str(exc),'expected':'successful native verification'})
    finally:
        if report.get('verifier_failed'):
            (logdir/'reward.txt').unlink(missing_ok=True);(logdir/'reward.json').unlink(missing_ok=True)
        else:(logdir/'reward.json').write_text(json.dumps({'reward':report.get('reward',0)},ensure_ascii=False,indent=2))
        (logdir/'criterion-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str))
        print(json.dumps({'status':report['status'],'reward':report.get('reward',0),'failed_checks':report.get('failed_checks',[]),'error':report.get('error'),'full_professional_score':None},ensure_ascii=False))
    return 0 if report['status']=='passed' else 1
if __name__=='__main__':raise SystemExit(main())
