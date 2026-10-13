#!/usr/bin/env python3
"""Acquire an explicit public-source plan; preserve raw bytes and parse evidence.

No login, discovery crawler, or benchmark authoring is performed. Run with a
Python environment containing pypdf and openpyxl for PDF/XLSX inspection.
"""
from __future__ import annotations

import argparse
import concurrent.futures
import hashlib
import io
import json
import re
import time
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]


class BodyText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.skip = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style'}:
            self.skip += 1
        if tag in {'p', 'div', 'br', 'tr', 'h1', 'h2', 'li'}:
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in {'script', 'style'} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def inspect(data: bytes, fmt: str):
    """Format validation and extraction, not a claim of business completeness."""
    if fmt == 'pdf':
        from pypdf import PdfReader
        if not data.startswith(b'%PDF-'):
            raise ValueError('Not a PDF response')
        reader = PdfReader(io.BytesIO(data), strict=False)
        pages = [p.extract_text() or '' for p in reader.pages]
        if not pages:
            raise ValueError('Empty PDF')
        return {'pages': len(pages), 'page_text_lengths': [len(t.strip()) for t in pages],
                'pages_needing_visual_review': [i + 1 for i, t in enumerate(pages) if len(t.strip()) < 40]}, pages
    if fmt == 'docx':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if z.testzip() is not None:
                raise ValueError('DOCX CRC failure')
            xml = ET.fromstring(z.read('word/document.xml'))
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            paragraphs = [''.join(n.text or '' for n in p.findall('.//w:t', ns))
                          for p in xml.findall('.//w:p', ns)]
            text = '\n'.join(paragraphs)
            return {'paragraph_text_nodes': len(paragraphs), 'tables': len(xml.findall('.//w:tbl', ns))}, [text]
    if fmt == 'xlsx':
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if z.testzip() is not None:
                raise ValueError('XLSX CRC failure')
        wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        sheets, text = [], []
        for ws in wb:
            cells = []
            for row in ws.iter_rows(values_only=True):
                values = [str(c) if c is not None else '' for c in row]
                while values and values[-1] == '':
                    values.pop()
                if values:
                    cells.append(values)
            text.append(ws.title + '\n' + '\n'.join('\t'.join(row) for row in cells))
            sheets.append({'name': ws.title, 'rows': ws.max_row, 'columns': ws.max_column})
        wb.close()
        return {'sheets': sheets}, text
    if fmt == 'zip':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if z.testzip() is not None:
                raise ValueError('ZIP CRC failure')
            members = [{'name': n.filename, 'bytes': n.file_size} for n in z.infolist()]
            pdfs = [n for n in z.namelist() if n.lower().endswith('.pdf')]
            pages = []
            for name in pdfs:
                info, member_pages = inspect(z.read(name), 'pdf')
                pages.extend([f'ZIP member: {name}; PDF page {i+1}\n{t}' for i, t in enumerate(member_pages)])
            return {'archive_members': members, 'pdf_members': pdfs}, pages
    if fmt == 'html':
        text = data.decode('utf-8', errors='replace')
        parser = BodyText()
        parser.feed(text)
        text = re.sub(r'\n\s*\n+', '\n\n', ''.join(parser.parts)).strip()
        return {'text_characters': len(text)}, [text]
    raise ValueError(f'Unsupported format: {fmt}')


def acquire(item, library):
    record = dict(item)
    for key in ('source_id', 'document_id'):
        if not re.fullmatch(r'[A-Za-z0-9_-]+', item[key]):
            raise ValueError(f'Unsafe path identifier: {key}')
    if urllib.parse.urlparse(item['original_url']).scheme not in {'http', 'https'}:
        raise ValueError('Only public HTTP(S) URLs are supported')
    record['acquired_at'] = datetime.now(ZoneInfo('Asia/Shanghai')).isoformat(timespec='seconds')
    if item.get('acquisition_policy') == 'link_only':
        record.update(acquisition_status='link_only', file_path=None, content_hash=None)
        return record
    folder = library / 'raw' / item['source_id'] / item['document_id']
    target = folder / ('original.' + item['original_format'])
    folder.mkdir(parents=True, exist_ok=True)
    try:
        metadata = {}
        if target.exists():
            data = target.read_bytes()
            old = folder / 'source_manifest.json'
            if old.exists():
                previous = json.loads(old.read_text())
                if hashlib.sha256(data).hexdigest() != previous['content_hash']:
                    raise ValueError('Cached original differs from recorded hash')
                parsed = library / 'parsed' / item['document_id']
                if (parsed/'pages.json').exists() and (parsed/'text.txt').exists():
                    previous.update(item)
                    return previous
                metadata = previous.get('http', {})
                record['acquired_at'] = previous.get('acquired_at', record['acquired_at'])
        else:
            url = urllib.parse.quote(item['original_url'], safe=':/?=&%+#@')
            last_error = None
            for attempt in range(2):
                try:
                    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0', 'Referer': item['publication_url']})
                    with urllib.request.urlopen(req, timeout=45) as res:
                        data = res.read(60 * 1024 * 1024 + 1)
                        if len(data) > 60 * 1024 * 1024:
                            raise ValueError('Response exceeds 60 MiB acquisition budget')
                        metadata = {'status': res.status, 'content_type': res.headers.get('Content-Type'),
                                    'final_url': res.geturl(), 'content_length': res.headers.get('Content-Length')}
                    break
                except Exception as exc:
                    last_error = exc
                    if attempt == 0:
                        time.sleep(1)
            else:
                raise last_error
        info, pages = inspect(data, item['original_format'])
        if item.get('expected_text') and pages:
            normalized = re.sub(r'\s+', '', '\n'.join(pages))
            missing = []
            for text in item['expected_text']:
                if re.sub(r'\s+', '', text) not in normalized:
                    missing.append(text)
            if missing:
                if not item.get('allow_visual_identity_review'):
                    raise ValueError(f'Expected identity/title missing: {missing}')
                info['identity_text_check'] = 'pending_visual_review'
                info['identity_text_missing'] = missing
        if not target.exists():
            target.write_bytes(data)
        parsed = library / 'parsed' / item['document_id']
        parsed.mkdir(parents=True, exist_ok=True)
        page_records = [{'page': i+1, 'text': t} for i, t in enumerate(pages)]
        (parsed/'pages.json').write_text(json.dumps(page_records, ensure_ascii=False, indent=2)+'\n')
        (parsed/'text.txt').write_text('\n\n'.join(f'[Page {i+1}]\n{t}' for i,t in enumerate(pages))+'\n')
        record.update(acquisition_status='downloaded', file_path=str(target.relative_to(ROOT)),
                      parsed_path=str(parsed.relative_to(ROOT)), content_hash=hashlib.sha256(data).hexdigest(),
                      bytes=len(data), http=metadata, format_check=info)
        (folder/'source_manifest.json').write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n')
    except Exception as exc:
        record.update(acquisition_status='failed', error=f'{type(exc).__name__}: {exc}', file_path=None, content_hash=None)
    return record


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('plan', type=Path)
    parser.add_argument('--workers', type=int, default=3)
    args = parser.parse_args()
    plan = json.loads(args.plan.read_text())
    library = ROOT / 'reference-library'
    records = []
    with concurrent.futures.ThreadPoolExecutor(max_workers=min(max(args.workers,1),4)) as pool:
        futures = {pool.submit(acquire, item, library): item for item in plan['documents']}
        for future in concurrent.futures.as_completed(futures):
            record = future.result()
            records.append(record)
            print(json.dumps({k:record.get(k) for k in ['document_id','acquisition_status','bytes','error']},ensure_ascii=False), flush=True)
    records.sort(key=lambda r:r['document_id'])
    (args.plan.parent/'catalog.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in records))
    return int(any(r['acquisition_status']=='failed' for r in records))


if __name__ == '__main__':
    raise SystemExit(main())
