#!/usr/bin/env python3
import csv
import re
from decimal import Decimal
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
import subprocess

ROOT = Path('/workspace')
INPUT = ROOT / 'input-data/CN-1225248741_original.pdf'
OUT = ROOT / 'output'
SOURCE = 'CN-1225248741_original.pdf'
PDF = PdfReader(INPUT)
PAGES = [(p.extract_text() or '') for p in PDF.pages]

# Page indexes here are physical PDF page numbers, not the printed folio.
FACTS = [
    (24, 'operating_inflow', '2025', '经营活动现金流入小计', '3,527,194,370.04'),
    (24, 'operating_inflow', '2024', '经营活动现金流入小计', '4,134,489,117.69'),
    (24, 'operating_outflow', '2025', '经营活动现金流出小计', '3,484,208,109.96'),
    (24, 'operating_outflow', '2024', '经营活动现金流出小计', '4,125,411,208.51'),
    (24, 'operating_net', '2025', '经营活动产生的现金流量净额', '42,986,260.08'),
    (24, 'operating_net', '2024', '经营活动产生的现金流量净额', '9,077,909.18'),
    (24, 'investing_net', '2025', '投资活动产生的现金流量净额', '99,170,943.48'),
    (24, 'investing_net', '2024', '投资活动产生的现金流量净额', '82,256,670.25'),
    (24, 'financing_inflow', '2025', '筹资活动现金流入小计', '974,280,000.00'),
    (24, 'financing_inflow', '2024', '筹资活动现金流入小计', '914,000,000.00'),
    (24, 'financing_outflow', '2025', '筹资活动现金流出小计', '1,100,823,838.67'),
    (24, 'financing_outflow', '2024', '筹资活动现金流出小计', '942,903,026.96'),
    (24, 'financing_net', '2025', '筹资活动产生的现金流量净额', '-126,543,838.67'),
    (24, 'financing_net', '2024', '筹资活动产生的现金流量净额', '-28,903,026.96'),
    (24, 'cash_change', '2025', '现金及现金等价物净增加额', '15,613,364.89'),
    (24, 'cash_change', '2024', '现金及现金等价物净增加额', '62,431,552.47'),
    (184, 'operating_receivables_change', '2025', '经营性应收项目的减少（增加以“－”号填列）', '33,779,318.53'),
    (184, 'operating_receivables_change', '2024', '经营性应收项目的减少（增加以“－”号填列）', '-30,000,712.68'),
    (184, 'supplement_operating_net', '2025', '经营活动产生的现金流量净额', '42,986,260.08'),
    (184, 'supplement_operating_net', '2024', '经营活动产生的现金流量净额', '9,077,909.18'),
]
GAPS = [
    ('G01', '本期原始试算平衡表、科目映射、现金流量表编制底稿及相关银行/票据凭证', '公开年报是汇总披露，不能建立逐科目、逐凭证错报金额或调整分录。', '错报汇总、借贷平衡及调整后试算平衡', '向项目档案管理员索取本期冻结版本及总账/明细账导出，核对主体范围和截止日。'),
    ('G02', '原始差异底稿、审计证据索引、差异原因及拟议调整记录', '仅凭年报列报数无法证明识别差异的完整性、来源和会计处理。', '错报完整性、分类与调整金额', '将差异逐项追溯至底稿和支持凭证，并与合并现金流量表编制过程核对。'),
    ('G03', '管理层接受/拒绝调整记录、过账凭证、过账后总账和客户书面确认', '审计软件或公开更正披露不等同于客户实际入账。', '调整状态及调整后余额交接', '逐笔追踪建议调整至经批准凭证、总账及更新后试算平衡；无法追踪者保留未确认状态。'),
    ('G04', '项目批准的重要性、执行重要性、明显微小错报阈值及适用报告基础', '公开年报不披露本项目的审计重要性，不能据此单独或汇总判断未更正错报。', '未更正错报单项及汇总评价、报告影响', '取得签批的审计计划和适用准则版本，再评价金额、性质、环境及以前期间影响。'),
    ('G05', '完整的已更正/未更正事项清单、管理层反馈、以前期间未更正错报滚存表', '单份年度报告没有项目层面的状态、反馈、期间滚存和沟通记录。', '未更正错报汇总、沟通及报告结论', '与错报汇总表和往年未更正事项表核对，留存管理层及治理层沟通证据。'),
]


def money(s):
    return Decimal(s.replace(',', ''))


def find_amount(page_num, label, expected):
    text = re.sub(r'\s+', '', PAGES[page_num - 1])
    compact_label = re.sub(r'\s+', '', label)
    if compact_label not in text:
        raise ValueError(f'Expected line item missing from PDF page {page_num}: {label}')
    if expected not in text:
        raise ValueError(f'Expected amount missing from PDF page {page_num}: {expected}')
    return expected


def write_csv(path, header, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def make_pdf(facts, checks):
    docx_path = OUT / 'misstatement_review.docx'
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = doc.styles['Normal']
    normal.font.name = 'Noto Sans CJK SC'
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(5)

    title = doc.add_heading('现金流量列报收尾复核备忘录', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('茂名石化实华股份有限公司（000637）｜公开材料复核｜金额：人民币元')
    doc.add_heading('一、范围与证据边界', level=2)
    doc.add_paragraph('本次仅检查输入文件 CN-1225248741_original.pdf，即2025年年度报告原件。范围为年报披露的2025年度及2024年度合并现金流量信息，引用均为PDF物理页。年报公开汇总数可用于列报复核和勾稽，不是项目原始试算平衡表、客户账簿或审计差异底稿。当前资料不支持形成审计调整分录、认定客户已入账、列示未更正错报余额或给出重大性结论。')
    doc.add_heading('二、公开披露金额及复核', level=2)
    preferred = ['经营活动现金流入小计','经营活动现金流出小计','经营活动产生的现金流量净额','筹资活动现金流入小计','筹资活动现金流出小计','筹资活动产生的现金流量净额','现金及现金等价物净增加额','经营性应收项目的减少（增加以“－”号填列）']
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['项目', '2025', '2024', '来源']
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for label in preferred:
        vals = {x[1]: x[4] for x in facts if x[3] == label}
        cells = table.add_row().cells
        for cell, value in zip(cells, [label, vals.get('2025','—'), vals.get('2024','—'), 'PDF第24页' if label != preferred[-1] else 'PDF第184页']):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        row.cells[0].width = Inches(3.0)
        row.cells[1].width = Inches(1.35)
        row.cells[2].width = Inches(1.35)
        row.cells[3].width = Inches(0.85)
    for title_text, y, result, note in checks:
        doc.add_paragraph(f'{title_text}（{y}）：{result}。{note}')
    doc.add_heading('三、限制与待核实事项', level=2)
    doc.add_paragraph('公开年度报告只显示披露后的合并/母公司报表和附注。要汇总审计差异，尚需冻结版原始试算平衡、科目映射、现金流量表编制底稿、逐项差异记录及底层凭证；要确认客户是否过账，需取得批准凭证、过账后总账/试算平衡及客户确认；要评价未更正错报，需取得项目重要性、适用报告基础、完整错报汇总与以前期间滚存及管理层反馈。具体补件和受阻决策见 evidence_gaps.csv。')
    doc.add_heading('四、建议后续程序（按顺序）', level=2)
    steps = [
        '1. 锁定报告签发前项目版本、主体范围和截止日，取得合并与母公司试算平衡及现金流量表编制底稿。',
        '2. 将披露差异逐项映射至总账科目、票据/银行凭证和会计处理依据，查明期间归属、列报分类及相关金融资产终止确认判断的证据链。',
        '3. 对拟议调整检查借贷平衡、科目映射、重复项目和跨期影响；用调整前余额重算调整后试算平衡。',
        '4. 逐项核对管理层接受/拒绝状态、过账凭证和更新后账簿；审计调整本身不作为客户入账证据。',
        '5. 取得已批准重要性及完整未更正错报表，分别和汇总评价金额、性质、环境与以前期间影响，并记录管理层/治理层沟通及报告结论。',
        '6. 只有完成上述核对后，才生成供客户更新账簿的调整及余额交接文件；当前资料下相关状态应保留为待补证。',
    ]
    for step in steps:
        doc.add_paragraph(step)
    p = doc.add_paragraph('口径说明：空白项目仅在本备忘录现金流净额勾稽中按未披露金额为零处理；此计算只用于公开报表展示金额的算术核对，不代表已核实现金流构成。')
    p.runs[0].font.size = Pt(8)
    doc.save(docx_path)
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', str(OUT), str(docx_path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    pdf_path = OUT / 'misstatement_review.pdf'
    if not pdf_path.is_file() or pdf_path.stat().st_size < 1000:
        raise RuntimeError('LibreOffice did not produce a valid memo PDF')


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rows=[]
    amounts={}
    for page, unit, period, label, expected in FACTS:
        amount=find_amount(page,label,expected)
        amounts[(unit,period)]=money(amount)
        rows.append([f'F{len(rows)+1:02d}',period,'合并',label,amount,'人民币元',SOURCE,page,'公开披露事实','年度报告披露值；不是客户总账或审计调整金额。'])
    write_csv(OUT/'evidence_register.csv', ['item_id','period','scope','line_item','reported_amount','unit','source_file','pdf_page','evidence_status','limitation'], rows)
    write_csv(OUT/'evidence_gaps.csv', ['gap_id','missing_evidence','why_needed','decision_blocked','next_action'], GAPS)
    checks=[]
    for y in ['2025','2024']:
        op_in=amounts[('operating_inflow',y)]; op_out=amounts[('operating_outflow',y)]; op_net=amounts[('operating_net',y)]
        fin_in=amounts[('financing_inflow',y)]; fin_out=amounts[('financing_outflow',y)]; fin_net=amounts[('financing_net',y)]
        inv_net=amounts[('investing_net',y)]; cash=amounts[('cash_change',y)]
        op_calc=op_in-op_out
        checks.append(('经营活动净额勾稽',y,f'{op_in:,.2f} − {op_out:,.2f} = {op_calc:,.2f}，报告列示 {op_net:,.2f}', '算术差额为零；来源PDF第24页。'))
        checks.append(('现金净增加额勾稽',y,f'{op_net:,.2f} + {inv_net:,.2f} + {fin_net:,.2f} = {op_net+inv_net+fin_net:,.2f}，报告列示 {cash:,.2f}', '汇率影响栏空白，按零仅作披露数字勾稽，范围限定于公开列报。'))
        if op_calc != op_net or op_net+inv_net+fin_net != cash:
            raise ValueError(f'Cash-flow reconciliation mismatch for {y}')
    make_pdf(rows,checks)

if __name__ == '__main__':
    main()
