#!/usr/bin/env python3
import argparse
import json
import math
import re
import sys
import traceback
from pathlib import Path

import fitz
from openpyxl import load_workbook


def load_contract():
    p = Path(__file__).with_name("assertions.json")
    data = json.loads(p.read_text(encoding="utf-8"))
    if data.get("scale") != len(data.get("assertions", [])):
        raise RuntimeError("assertion scale mismatch")
    return data


def read_inputs(root):
    docs = {}
    for name in [
        "CN-1225327554.pdf", "CN-1225327559.pdf", "CN-1225327560.pdf",
        "CN-1225449371.pdf", "CN-1225449382.pdf", "CN-1225449388.pdf",
        "CN-1225449390.pdf",
    ]:
        p = root / name
        if not p.is_file() or p.stat().st_size < 1000:
            raise RuntimeError(f"trusted input missing or truncated: {name}")
        try:
            with fitz.open(p) as pdf:
                docs[name] = "\n".join(page.get_text() for page in pdf)
        except Exception as e:
            raise RuntimeError(f"trusted input cannot be parsed: {name}: {e}")
    normalized_580 = normalize(docs["CN-1225327560.pdf"])
    normalized_650 = normalize(docs["CN-1225449390.pdf"])
    if "236,800.00" not in docs["CN-1225449388.pdf"] or "2025年12月31日" not in normalized_580 or "2026年4月30日" not in normalized_650 or "2026年1-4月" not in normalized_650:
        raise RuntimeError("trusted input anchor absent")
    return docs


def candidate_rows(output):
    path = output / "dd_request_tracker.xlsx"
    if not path.is_file() or path.stat().st_size == 0:
        return [], "tracker missing"
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if "需求跟踪" not in wb.sheetnames:
            return [], "需求跟踪 sheet missing"
        ws = wb["需求跟踪"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], "tracker empty"
        headers = [str(x or "").strip() for x in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(v is not None for v in row)], None
    except Exception as e:
        return [], f"tracker parse error: {e}"


def normalize(v):
    return re.sub(r"\s+", "", str(v or "")).lower()


def assess(output, docs):
    rows, parse_issue = candidate_rows(output)
    by_id = {str(r.get("需求编号", "")).strip(): r for r in rows}
    summary_path = output / "dd_gap_summary.pdf"
    summary_text, summary_issue = "", None
    if not summary_path.is_file() or summary_path.stat().st_size == 0:
        summary_issue = "summary PDF missing"
    else:
        try:
            with fitz.open(summary_path) as pdf:
                summary_text = "\n".join(page.get_text() for page in pdf)
            if not summary_text.strip():
                summary_issue = "summary PDF has no extractable text"
        except Exception as e:
            summary_issue = f"summary PDF parse error: {e}"

    revised = docs["CN-1225449388.pdf"]
    old_audit = docs["CN-1225327560.pdf"]
    new_audit = docs["CN-1225449390.pdf"]
    note = docs["CN-1225449371.pdf"]
    checks = {}
    summary_normalized = normalize(summary_text)
    summary_requirements = [
        "236,800.00万元",
        "2025-12-31",
        "8-580",
        "2024、2025年度及年末",
        "8-650",
        "2026年1-4月",
        "扩界扩能手续仍在办理",
        "尚需股东会",
        "数据室索引",
    ]
    checks["PDF.summary-content"] = (
        not summary_issue and all(normalize(item) in summary_normalized for item in summary_requirements),
        {"extractable_text": bool(summary_normalized), "required_anchors": summary_requirements},
    )

    def has(row, fields):
        text = normalize(" ".join(str(row.get(f, "") or "") for f in fields))
        return bool(text)

    def rowcheck(rid, predicate):
        row = by_id.get(rid)
        if not row:
            return False, "missing row"
        try:
            return bool(predicate(row)), {k: row.get(k) for k in ["调查事项", "材料对应", "覆盖状态", "状态依据与限制", "追问/补件"]}
        except Exception as e:
            return False, str(e)

    def rowtext(r):
        return normalize(" ".join(str(v or "") for v in r.values()))

    checks["R01.scope"] = rowcheck("R01", lambda r: all(x in rowtext(r) for x in ["40%", "国城实业", "中信信托"]) and "cn-1225449388.pdf" in normalize(str(r.get("材料对应", ""))) and "236,800.00" in revised)
    checks["R02.version"] = rowcheck("R02", lambda r: all(x in rowtext(r) for x in ["cn-1225327554.pdf", "cn-1225449388.pdf", "cn-1225449371.pdf", "修订"]) and "修订说明" in note)
    checks["R03.price"] = rowcheck("R03", lambda r: all(x in rowtext(r) for x in ["2025-12-31", "236,800.00", "cn-1225449388.pdf", "592,070.01"]) and "236,800.00" in revised)
    checks["R04.audit-period"] = rowcheck("R04", lambda r: all(x in rowtext(r) for x in ["cn-1225327560.pdf", "2024-01-01", "2025-12-31"]) and "2026年1-4月" not in normalize(old_audit))
    checks["R05.audit-extension"] = rowcheck("R05", lambda r: all(x in rowtext(r) for x in ["cn-1225449390.pdf", "2026-04-30", "2026年1-4月"]) and "2026年1-4月" in normalize(new_audit))
    checks["R06.revision-topics"] = rowcheck("R06", lambda r: "cn-1225449371.pdf" in rowtext(r) and "修订" in rowtext(r) and all(x in normalize(note) for x in ["加期财务数据", "行政处罚", "采矿权抵押"]))
    checks["R07.legal-opinion"] = rowcheck("R07", lambda r: "cn-1225327559.pdf" in rowtext(r) and any(x in rowtext(r) for x in ["扫描", "机器文本", "页面"]))
    checks["R08.supplemental-opinion"] = rowcheck("R08", lambda r: "cn-1225449382.pdf" in rowtext(r) and any(x in rowtext(r) for x in ["扫描", "页面", "可检索"]))
    checks["R09.mining-rights"] = rowcheck("R09", lambda r: all(x in rowtext(r) for x in ["采矿", "探矿", "许可证"]) and "未提供" not in str(r.get("覆盖状态", "")))
    checks["R10.expansion"] = rowcheck("R10", lambda r: all(x in rowtext(r) for x in ["500", "800", "cn-1225449388.pdf"]) and any(x in rowtext(r) for x in ["仍在办理", "未完成"]) and "500万吨/年" in normalize(revised) and "800万吨/年" in normalize(revised))
    checks["R11.mortgage"] = rowcheck("R11", lambda r: all(x in rowtext(r) for x in ["抵押", "清偿", "cn-1225449388.pdf"]) and any(x in str(r.get("追问/补件", "")) for x in ["注销", "解除", "登记"]))
    checks["R12.equity-encumbrance"] = rowcheck("R12", lambda r: all(x in rowtext(r) for x in ["质押", "冻结", "cn-1225449388.pdf"]) and any(x in str(r.get("追问/补件", "")) for x in ["股东名册", "登记", "查询"]))
    checks["R13.penalties"] = rowcheck("R13", lambda r: "7" in rowtext(r) and "cn-1225449388.pdf" in rowtext(r) and any(x in str(r.get("追问/补件", "")) for x in ["决定书", "整改"]) and "受到行政处罚共计7起" in normalize(revised))
    checks["R14.related-party"] = rowcheck("R14", lambda r: "关联" in rowtext(r) and all(x in str(r.get("追问/补件", "")) for x in ["合同", "定价"]) and any(x in str(r.get("追问/补件", "")) for x in ["决策", "决议"]))
    checks["R15.approvals"] = rowcheck("R15", lambda r: all(x in rowtext(r) for x in ["尚需", "股东会", "cn-1225449388.pdf"]) and any(x in rowtext(r) for x in ["未完成", "跟踪", "未获批准"]) and not any(x in rowtext(r) for x in ["股东会已审议通过", "已完成全部程序", "程序已完成", "无需跟踪", "已独立核实"]))
    checks["R16.agreement"] = rowcheck("R16", lambda r: any(x in str(r.get("覆盖状态", "")) for x in ["未提供", "部分覆盖"]) and "协议" in str(r.get("追问/补件", "")) and any(x in str(r.get("追问/补件", "")) for x in ["交割", "先决条件"]))
    checks["R17.earnout"] = rowcheck("R17", lambda r: all(x in rowtext(r) for x in ["业绩承诺", "补偿", "cn-1225449388.pdf"]) and any(x in rowtext(r) for x in ["原件", "未提供"]))
    checks["R18.audit-integrity"] = rowcheck("R18", lambda r: "cn-1225449390.pdf" in rowtext(r) and any(x in str(r.get("追问/补件", "")) for x in ["签章", "验签"]) and any(x in str(r.get("追问/补件", "")) for x in ["附件", "完整"]))
    checks["R19.corporate-records"] = rowcheck("R19", lambda r: all(x in str(r.get("追问/补件", "")) for x in ["营业执照", "章程"]) and any(x in str(r.get("覆盖状态", "")) for x in ["未提供", "部分覆盖"]))
    checks["R20.ownership-chain"] = rowcheck("R20", lambda r: "股东名册" in str(r.get("追问/补件", "")) and any(x in str(r.get("追问/补件", "")) for x in ["出资", "股权变动"]) and "权属" in rowtext(r))
    checks["R21.contracts"] = rowcheck("R21", lambda r: "重大合同" in rowtext(r) and "合同台账" in str(r.get("追问/补件", "")) and any(x in str(r.get("追问/补件", "")) for x in ["违约", "补充协议"]))
    checks["R22.assets"] = rowcheck("R22", lambda r: all(x in str(r.get("追问/补件", "")) for x in ["权证", "在建工程"]) and "土地" in str(r.get("调查事项", "")))
    checks["R23.employment"] = rowcheck("R23", lambda r: all(x in str(r.get("追问/补件", "")) for x in ["员工名册", "社保"]) and "劳动" in str(r.get("调查事项", "")))
    checks["R24.data-room"] = rowcheck("R24", lambda r: all(x in rowtext(r) for x in ["数据室", "上轮"]) and any(x in str(r.get("追问/补件", "")) for x in ["上传时间", "版本"]))

    results = []
    for a in load_contract()["assertions"]:
        passed, actual = checks.get(a["id"], (False, "no checker"))
        results.append({"id": a["id"], "kind": a["kind"], "criterion": a["criterion"], "status": "passed" if passed else "failed", "actual": actual, "expected": "公开要求对应的项目化需求、材料锚点和可执行追问均可核验"})
    return results, [x for x in [parse_issue, summary_issue] if x]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input-data", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--logs", required=True)
    args = ap.parse_args()
    logs = Path(args.logs)
    logs.mkdir(parents=True, exist_ok=True)
    report_path, reward_path = logs / "report.json", logs / "reward.json"
    for p in (report_path, reward_path):
        try:
            p.unlink()
        except FileNotFoundError:
            pass
    try:
        docs = read_inputs(Path(args.input_data))
        assertions, defects = assess(Path(args.output), docs)
        contract = load_contract()
        passed = sum(x["status"] == "passed" for x in assertions)
        scale = contract["scale"]
        score = passed / scale
        caps = [ce["cap"] for ce in contract.get("critical_errors", []) if next((x for x in assertions if x["id"] == ce["assertion"]), {}).get("status") == "failed"]
        if caps:
            score = min(score, min(caps))
        status = "graded_candidate_defect" if defects else "graded"
        report = {"status": status, "valid": True, "reward": score, "P": passed, "S": scale, "assertions": assertions, "candidate_defects": defects, "critical_error_caps": caps}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        reward_path.write_text(json.dumps({"reward": score}, ensure_ascii=False) + "\n", encoding="utf-8")
    except Exception as e:
        report = {"status": "verifier_failure", "valid": False, "error": f"{type(e).__name__}: {e}", "traceback": traceback.format_exc()}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        for p in (reward_path,):
            try:
                p.unlink()
            except FileNotFoundError:
                pass
        print(report["error"], file=sys.stderr)
        sys.exit(2)

if __name__ == "__main__":
    main()
