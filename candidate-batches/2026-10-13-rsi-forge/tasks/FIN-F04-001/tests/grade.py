#!/usr/bin/env python3
import json, math, os, re, sys
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader

OUT=Path('/workspace/output'); LOG=Path('/logs/verifier'); LOG.mkdir(parents=True,exist_ok=True)
ids=[
'scope.next_month_august','scope.single_bond','bond.code','bond.short_name','bond.issuer','bond.balance','bond.maturity','bond.payment_node','bond.rate','amount.interest','amount.total','amount.formula','funding.no_proof','funding.confirmed_zero','funding.unconfirmed','funding.gap','funding.dedup','evidence.trustee_page','evidence.notice_rule','evidence.annual_cash_limit','risk.financial','risk.credit','risk.enhancement','actions.priority','artifact.editable']
results=[]

def add(i,ok,actual,expected): results.append({'id':i,'status':'passed' if ok else 'failed','actual':actual,'expected':expected})
try:
 x=OUT/'debt_funding_check.xlsx'; p=OUT/'funding_risk_note.pdf'
 wb=load_workbook(x,data_only=False); wb_values=load_workbook(x,data_only=True); ws=wb['排查表']; calc=wb['计算与证据']; calc_values=wb_values['计算与证据']
 vals={ws.cell(4,c).value for c in range(1,18)}
 row=[ws.cell(4,c).value for c in range(1,18)]
 pdf='\n'.join(pg.extract_text() or '' for pg in PdfReader(str(p)).pages)
 calcrow={calc.cell(1,c).value:calc.cell(2,c).value for c in range(1,16)}
 calcrow_values={calc_values.cell(1,c).value:calc_values.cell(2,c).value for c in range(1,16)}
 formula_d=calc['D2'].value; formula_e=calc['E2'].value
 add('scope.next_month_august', '2026年8月' in pdf and '排查截止日：2026年7月31日' in pdf and '2026-08' in str(calcrow.get('合同到期/回售日')), 'scope', '2026年8月/2026-07-31')
 add('scope.single_bond', row[2]=='21申宏06' and ws.max_row>=4, '21申宏06', '唯一逐券记录')
 add('bond.code', row[1]=='149579.SZ', row[1], '149579.SZ')
 add('bond.short_name', row[2]=='21申宏06', row[2], '21申宏06')
 add('bond.issuer', row[3]=='申万宏源集团股份有限公司' and '申万宏源集团股份有限公司' in pdf, row[3], '申万宏源集团股份有限公司')
 add('bond.balance', abs(float(row[7])-25)<1e-9 and '25.00' in pdf, row[7], 25.0)
 add('bond.maturity', row[9]=='2026-08-04' or str(row[9])[:10]=='2026-08-04' or '2026-08-04' in pdf, row[9], '2026-08-04')
 add('bond.payment_node', str(calcrow.get('实际付款节点'))[:10]=='2026-08-04' and '实际付款节点' in pdf, calcrow.get('实际付款节点'), '2026-08-04')
 add('bond.rate', abs(float(calcrow.get('票面利率'))-0.0338)<1e-8 and '3.38%' in pdf, calcrow.get('票面利率'), 0.0338)
 add('amount.interest', '0.845' in pdf and '0.85' in pdf and abs(float(calcrow_values.get('当期利息（亿元）') or 0)-0.85)<1e-9, calcrow_values.get('当期利息（亿元）'), 0.85)
 add('amount.total', abs(float(calcrow_values.get('本金及利息合计（亿元）') or 0)-25.85)<1e-9 and '25.85' in pdf, calcrow_values.get('本金及利息合计（亿元）'), 25.85)
 add('amount.formula', isinstance(formula_d,str) and formula_d.startswith('=') and isinstance(formula_e,str) and formula_e.startswith('='), [formula_d,formula_e], 'formula cells')
 add('funding.no_proof', '专项资金证明' in pdf and ('不能作为已落实' in pdf or '不等于本次25.85亿元已划入' in pdf), 'proof distinction', 'explicit')
 add('funding.confirmed_zero', '已取得、可用且有到账时点证据的资金：0.00亿元' in pdf, '0.00', '0.00')
 add('funding.unconfirmed', '待确认/落实中的计划性来源：25.85亿元' in pdf and '不计入已落实' in pdf, '25.85 unconfirmed', '25.85 unconfirmed')
 add('funding.gap', '可核实缺口：25.85亿元' in pdf, '25.85', '25.85')
 add('funding.dedup', '去重' in pdf and '去重核对' in str(calcrow), 'dedup recorded', 'dedup recorded')
 add('evidence.trustee_page', '02_申万宏源2025年度受托报告.pdf第8页' in pdf and '原件页码7' in pdf, 'trustee p8', 'trustee p8')
 add('evidence.notice_rule', '新疆证监局' in pdf and '二（三）' in pdf and '二（四）' in pdf, 'notice clauses', 'clauses 2(3), 2(4)')
 add('evidence.annual_cash_limit', '1,883.91亿元' in pdf and '不等于本次25.85亿元已划入' in pdf, 'cash limitation', 'explicit')
 add('risk.financial', '1,883.91亿元' in pdf and '1,231.60亿元' in pdf, 'financial context', 'cash and bond liabilities')
 add('risk.credit', '未见受托报告所述违约' in str(row[12]) and '重大事项' in pdf, 'credit observation', 'source-grounded')
 add('risk.enhancement', '无担保' in pdf and 'AAA' in pdf and '局限' in pdf, 'enhancement/rating', 'limits stated')
 add('actions.priority', '优先级1' in pdf and '电话报告' in pdf and '2个工作日' in pdf and '逐日' in pdf, 'prioritized actions', 'actions')
 headers=[ws.cell(3,c).value for c in range(1,18)]
 req=['证券代码','证券简称','发行人中文名称','债券余额（亿元）','行权日期','日期性质','风险分类','兑付资金落实情况','兑付资金来源','备注']
 add('artifact.editable', all(h in headers for h in req) and bool(ws.auto_filter.ref) and '计算与证据' in wb.sheetnames, headers, 'required fields/filter/evidence sheet')
except Exception as e:
 report={'reward':None,'P':0,'S':len(ids),'valid':False,'status':'verifier_failure','error':repr(e),'assertions':results}
 (LOG/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8'); raise
passed=sum(r['status']=='passed' for r in results); S=len(results); reward=passed/S
report={'reward':reward,'P':passed,'S':S,'valid':True,'status':'graded' if passed else 'graded_candidate_defect','assertions':results}
(LOG/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
(LOG/'reward.json').write_text(json.dumps({'reward':reward}),encoding='utf-8')
print(json.dumps({'reward':reward,'P':passed,'S':S}))
