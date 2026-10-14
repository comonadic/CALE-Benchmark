#!/usr/bin/env python3
import argparse
import hashlib
import json
import math
import re
import subprocess
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

EXPECTED = {
    'COURT-ZHENXIN-CLAIM.docx': 'd818a63eb2a4f7208cda898e4863e3bcea7e61b999e8fc60a28ebb4bc75a17d5',
    'COURT-ZHENXIN-GUIDE.pdf': 'aeecb43d1d96cdf1ac02e0baa2ad872c40edaa22b6e4a58cde7d5893ed478a7b',
    'COURT-ZHENXIN-LIST.docx': '6d337bbe0809738406489588245be8b7cc089c9d4d82d13f4b1b54665725fc94',
    'COURT-ZHENXIN-REGISTER.docx': 'a0941539d4e66027e585997801883b31b5fff04019c085fab14eea048fea67db',
}
ASSERTIONS_PATH = Path(__file__).with_name('assertions.json')


def normalize(s):
    return re.sub(r'\s+', '', s or '')


def read_docs(input_dir):
    docs = {}
    for filename in ('COURT-ZHENXIN-CLAIM.docx', 'COURT-ZHENXIN-LIST.docx', 'COURT-ZHENXIN-REGISTER.docx'):
        doc = Document(input_dir / filename)
        text = '\n'.join(p.text for p in doc.paragraphs)
        text += '\n' + '\n'.join(c.text for t in doc.tables for row in t.rows for c in row.cells)
        docs[filename] = text
    return docs


def pdf_text(path):
    reader = PdfReader(path)
    return '\n'.join(page.extract_text() or '' for page in reader.pages)


def check_formula_recalculation(xlsx):
    with tempfile.TemporaryDirectory(prefix='lawl04-lo-') as tmp:
        profile = Path(tmp) / 'profile'
        out = Path(tmp) / 'out'
        out.mkdir()
        run = subprocess.run(['libreoffice', f'-env:UserInstallation=file://{profile}', '--headless',
                              '--convert-to', 'xlsx', '--outdir', str(out), str(xlsx)],
                             capture_output=True, text=True, timeout=90)
        recalced = out / xlsx.name
        if run.returncode != 0 or not recalced.exists():
            raise RuntimeError('LibreOffice recalc failed: ' + run.stderr[-500:])
        cached = load_workbook(recalced, data_only=True)
        ws = cached['计息明细']
        days, interest = ws['G2'].value, ws['H2'].value
        if days not in (None, '') or interest not in (None, ''):
            raise RuntimeError(f'Blank-input formulas must remain blank, got {days!r}, {interest!r}')
        # Test the actual formula dependency on an isolated copy, then recalculate it.
        test_input = Path(tmp) / 'formula-test.xlsx'
        formula_wb = load_workbook(xlsx)
        formula_ws = formula_wb['计息明细']
        formula_ws['C2'] = date(2026, 8, 1)
        formula_ws['D2'] = date(2026, 9, 2)
        formula_ws['E2'] = 10000
        formula_ws['F2'] = 0.12
        formula_wb.save(test_input)
        test_out = Path(tmp) / 'test-out'
        test_out.mkdir()
        test_profile = Path(tmp) / 'profile2'
        trial = subprocess.run(['libreoffice', f'-env:UserInstallation=file://{test_profile}', '--headless',
                                '--convert-to', 'xlsx', '--outdir', str(test_out), str(test_input)],
                               capture_output=True, text=True, timeout=90)
        result_path = test_out / test_input.name
        if trial.returncode != 0 or not result_path.exists():
            raise RuntimeError('LibreOffice formula mutation recalc failed')
        values = load_workbook(result_path, data_only=True)['计息明细']
        if values['G2'].value != 32 or not math.isclose(values['H2'].value or 0, 106.67, abs_tol=0.01):
            raise RuntimeError(f'Formula dependency mismatch: days={values["G2"].value}, interest={values["H2"].value}')


def grade(input_dir, output_dir):
    hashes = {name: hashlib.sha256((input_dir / name).read_bytes()).hexdigest() for name in EXPECTED}
    if hashes != EXPECTED:
        raise RuntimeError('Trusted input hash mismatch')
    docs = read_docs(input_dir)
    guide = pdf_text(input_dir / 'COURT-ZHENXIN-GUIDE.pdf')
    norm_guide = normalize(guide)
    claim_template = normalize(docs['COURT-ZHENXIN-CLAIM.docx'])
    register_template = normalize(docs['COURT-ZHENXIN-REGISTER.docx'])
    list_template = normalize(docs['COURT-ZHENXIN-LIST.docx'])

    xlsx = output_dir / 'claim_calculation.xlsx'
    decl = output_dir / 'claim_declaration.pdf'
    index = output_dir / 'supporting_evidence_index.pdf'
    missing = [p.name for p in (xlsx, decl, index) if not p.is_file() or p.stat().st_size == 0]
    files_present = not missing
    assertions_doc = json.loads(ASSERTIONS_PATH.read_text())
    if not files_present:
        return {a['id']: (False, {'missing_files': missing}) for a in assertions_doc['assertions']}

    wb = load_workbook(xlsx, data_only=False)
    sheet_names = set(wb.sheetnames)
    workbook_valid = {'事项与口径', '计息明细', '补件清单'}.issubset(sheet_names)
    calc = wb['计息明细'] if '计息明细' in sheet_names else None
    matter = wb['事项与口径'] if '事项与口径' in sheet_names else None
    gaps = wb['补件清单'] if '补件清单' in sheet_names else None
    matter_text = normalize(' '.join(str(c.value or '') for row in matter.iter_rows() for c in row)) if matter else ''
    calc_text = normalize(' '.join(str(c.value or '') for row in calc.iter_rows() for c in row)) if calc else ''
    gap_text = normalize(' '.join(str(c.value or '') for row in gaps.iter_rows() for c in row)) if gaps else ''
    decl_text = normalize(pdf_text(decl))
    index_text = normalize(pdf_text(index))
    decl_doc = Document(input_dir / 'COURT-ZHENXIN-CLAIM.docx')
    template_claim_text = normalize('\n'.join(p.text for p in decl_doc.paragraphs) + '\n' +
                                    '\n'.join(c.text for t in decl_doc.tables for r in t.rows for c in r.cells))

    def has(text, *terms):
        return all(normalize(term) in text for term in terms)

    def outcome(aid, ok, actual, expected):
        return (bool(ok), {'actual': actual, 'expected': expected})

    results = {}
    received_cell = matter['B3'].value if matter and matter.max_row >= 3 else None
    deadline_cell = matter['B5'].value if matter and matter.max_row >= 5 else None
    received_date = received_cell.date() if hasattr(received_cell, 'date') else received_cell
    deadline_date = deadline_cell.date() if hasattr(deadline_cell, 'date') else deadline_cell
    results['PROC.receipt-date'] = outcome('PROC.receipt-date', received_date == date(2026, 9, 2) and '2026年9月2日' in decl_text,
                                           str(received_cell), '2026-09-02; declaration identifies this cutoff')
    results['PROC.deadline'] = outcome('PROC.deadline', deadline_date == date(2026, 11, 3) and '2026年11月3日' in index_text,
                                       str(deadline_cell), '2026-11-03 in workbook and evidence index')
    results['PROC.meeting'] = outcome('PROC.meeting', has(matter_text, '2026年11月16日10时') and '2026年11月16日10时' in normalize(guide),
                                      'meeting entry', '2026-11-16 10:00 with guide source')
    results['PROC.paper-electronic'] = outcome('PROC.paper-electronic', has(matter_text, '盖章扫描件', '邮寄纸质材料'),
                                               'submission procedure text', 'scan and paper-mail steps distinguished')
    results['PROC.paper-receipt'] = outcome('PROC.paper-receipt', has(matter_text, '管理人收到的纸质材料', '时间和内容') or has(matter_text, '管理人收到纸质材料的时间和内容'),
                                            'receipt rule', 'paper receipt controls content and time')
    ordered_terms = ['债权申报材料清单', '债权申报登记表', '债权申报书', '主体及授权材料', '送达方式确认书', '债权事实证据']
    order_variants = [('债权申报材料清单', '材料清单'), ('债权申报登记表', '登记表'),
                      ('债权申报书', '申报书'), ('主体及授权材料', '主体及授权材料'),
                      ('送达方式确认书', '送达方式确认书'), ('债权事实证据', '债权事实证据')]
    order_positions = [next((matter_text.find(term) for term in variants if matter_text.find(term) >= 0), -1)
                       for variants in order_variants]
    order_text = all(pos >= 0 for pos in order_positions) and order_positions == sorted(order_positions)
    results['PROC.order'] = outcome('PROC.order', order_text, 'workbook material sequence', ordered_terms)
    results['PROC.count'] = outcome('PROC.count', has(matter_text, '一式一份', '邮寄底单'),
                                    'paper count and proof of dispatch', 'one paper set; retain postal receipt')

    cutoff_cell = matter['B4'].value if matter and matter.max_row >= 4 else None
    cutoff_date = cutoff_cell.date() if hasattr(cutoff_cell, 'date') else cutoff_cell
    row_cutoff = calc['D2'].value if calc and calc.max_row >= 2 else None
    row_cutoff_date = row_cutoff.date() if hasattr(row_cutoff, 'date') else row_cutoff
    results['CALC.exclusive-cutoff'] = outcome('CALC.exclusive-cutoff', cutoff_date == date(2026, 9, 2) and row_cutoff_date == date(2026, 9, 2) and
                                               has(cutoff_text, '2026-09-02', '不含') or
                                               (cutoff_date == date(2026, 9, 2) and row_cutoff_date == date(2026, 9, 2) and has(cutoff_text, '2026年9月2日', '不含当天')),
                                               {'workbook_cutoff': str(cutoff_cell), 'line_cutoff': str(row_cutoff)}, 'workbook and calculation row cutoff 2026-09-02, exclusive')
    formula_ok = bool(calc and isinstance(calc['G2'].value, str) and calc['G2'].value.startswith('=IF(') and
                      isinstance(calc['H2'].value, str) and 'E2*F2*G2/360' in calc['H2'].value)
    results['CALC.day-count'] = outcome('CALC.day-count', formula_ok and 'D2-C2' in calc['G2'].value,
                                        {'G2': calc['G2'].value if calc else None}, 'actual date difference formula with blank guard')
    formula_recalc_ok = False
    if formula_ok:
        check_formula_recalculation(xlsx)
        formula_recalc_ok = True
    results['CALC.interest-formula'] = outcome('CALC.interest-formula', formula_ok and formula_recalc_ok,
                                               {'formula': calc['H2'].value if calc else None, 'mutation_recalc': formula_recalc_ok},
                                               '360-day simple interest formula recalculates to 32 days / CNY 106.67 on isolated probe')
    results['CALC.blank-not-zero'] = outcome('CALC.blank-not-zero', bool(calc and calc['E2'].value is None and calc['F2'].value is None and
                                                                          calc['I2'].value == '待补' and
                                                                          has(decl_text, '待委托方债权凭证', '尚待核验')),
                                              {'principal': calc['E2'].value if calc else None, 'status': calc['I2'].value if calc else None},
                                              'no fabricated amount; pending source evidence')
    results['CALC.source-anchor'] = outcome('CALC.source-anchor', has(matter_text, 'COURT-ZHENXIN-GUIDE.pdf', '第4页', '360天'),
                                            'workbook source locator', 'guide page 4 and 360-day convention')

    results['GAP.corporate-authority'] = outcome('GAP.corporate-authority', has(gap_text, '营业执照', '授权委托书', '律所公函'),
                                                  'corporate and authorization evidence rows', 'business identity and representative authority evidence listed')
    results['GAP.underlying'] = outcome('GAP.underlying', has(gap_text, '合同/协议', '交付/验收', '对账'),
                                         'claim origin evidence rows', 'contract and performance/ledger evidence listed')
    results['GAP.repayments-offsets'] = outcome('GAP.repayments-offsets', has(gap_text, '收款、还款、抵销', '避免重复申报'),
                                                 'payment/offset row', 'payments and offsets reconciled to avoid duplicate claim')
    results['GAP.security'] = outcome('GAP.security', has(gap_text, '担保合同', '登记证明', '担保范围'),
                                       'security evidence row', 'security status and classification await evidence')
    results['GAP.proceedings'] = outcome('GAP.proceedings', has(gap_text, '仲裁文书', '执行受理', '执行到款'),
                                          'proceedings evidence row', 'litigation/arbitration/enforcement and receipts listed')

    fields = ['申报人', '统一社会信用代码/身份证号', '住所', '本金', '孳息/违约金/滞纳金', '诉讼费用', '债权性质', '事实和理由', '签字/盖章']
    field_coverage = all(normalize(x) in template_claim_text for x in fields)
    results['DOC.declaration-fields'] = outcome('DOC.declaration-fields', field_coverage and has(decl_text, '申报人', '统一社会信用代码', '债权性质', '事实和理由', '签字/盖章'),
                                                 'fields preserved from blank filing template', fields)
    results['DOC.declaration-status'] = outcome('DOC.declaration-status', has(decl_text, '待补资料草稿', '未签署', '尚待核验', '不主张特定债权金额'),
                                                 'draft and pending language', 'not signed/submitted; no unsupported claim assertion')
    results['DOC.declaration-crossref'] = outcome('DOC.declaration-crossref', has(decl_text, '东莞市震鑫磁业有限公司', '2026年9月2日', '2026年11月3日') and
                                                  '2026年11月3日' in index_text and has(matter_text, '2026-11-03'),
                                                   'debtor/date cross-reference', 'case debtor and controlling dates match across artifacts')

    order_cover = all(term in index_text for term in ['债权申报材料清单', '债权申报登记表', '债权申报书', '主体资料及授权材料', '送达方式确认书', '债权形成及履行证据'])
    results['IDX.order'] = outcome('IDX.order', order_cover, 'evidence index categories', ordered_terms)
    results['IDX.status-separation'] = outcome('IDX.status-separation', has(index_text, '管理人空白原件已提供', '待补', '不是债权事实证据'),
                                               'provided source vs missing creditor materials', 'status and evidentiary role separated')
    results['IDX.no-false-receipt'] = outcome('IDX.no-false-receipt', has(index_text, '不是实际提交/签收清单', '待收件后填写') and
                                              not has(index_text, '已签收', '已提交管理人'),
                                              'filing/receipt status', 'index is work draft, no false filing or receipt claim')
    input_ok = all(hashes[k] == EXPECTED[k] for k in EXPECTED)
    results['PKG.originals'] = outcome('PKG.originals', input_ok and all(x in guide for x in ['东莞市震鑫磁业有限公司', '2026 年 11 月 3 日前']) and
                                       '债权申报书' in claim_template and '债权申报登记表' in register_template and '债权申报材料清单' in list_template,
                                       hashes, 'four public originals match indexed SHA-256 and are parseable')
    return results


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input-data', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--logs', type=Path, required=True)
    args = parser.parse_args()
    report_path = args.logs / 'report.json'
    reward_path = args.logs / 'reward.json'
    report_path.parent.mkdir(parents=True, exist_ok=True)
    for stale in (report_path, reward_path):
        if stale.exists(): stale.unlink()
    contract = json.loads(ASSERTIONS_PATH.read_text())
    try:
        results = grade(args.input_data, args.output)
        rows = []
        for assertion in contract['assertions']:
            passed, detail = results[assertion['id']]
            rows.append({'id': assertion['id'], 'status': 'passed' if passed else 'failed',
                         'actual': detail.get('actual'), 'expected': detail.get('expected')})
        valid = True
        p = sum(row['status'] == 'passed' for row in rows)
        missing_outputs = any(not (args.output / name).is_file() or (args.output / name).stat().st_size == 0
                              for name in ('claim_calculation.xlsx', 'claim_declaration.pdf', 'supporting_evidence_index.pdf'))
        status = 'graded_candidate_defect' if missing_outputs else 'graded'
        ratio = min(1.0, p / contract['scale'])
        caps = []
        by_id = {row['id']: row for row in rows}
        for ce in contract.get('critical_errors', []):
            if by_id[ce['assertion']]['status'] == 'failed':
                ratio = min(ratio, ce['cap'])
                caps.append({'id': ce['id'], 'assertion': ce['assertion'], 'cap': ce['cap']})
        report = {'reward': ratio, 'P': p, 'S': contract['scale'], 'valid': valid, 'status': status,
                  'assertions': rows, 'critical_error_caps': caps}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
        reward_path.write_text(json.dumps({'reward': ratio}, ensure_ascii=False))
    except Exception as exc:
        report = {'valid': False, 'status': 'verifier_failure', 'error': f'{type(exc).__name__}: {exc}',
                  'assertions': [], 'reward': None, 'P': None, 'S': contract['scale']}
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2))
        if reward_path.exists(): reward_path.unlink()
        raise


if __name__ == '__main__':
    main()
