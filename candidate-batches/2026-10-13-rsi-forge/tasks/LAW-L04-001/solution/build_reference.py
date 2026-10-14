#!/usr/bin/env python3
import argparse
from datetime import date
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                Spacer, Table, TableStyle)

RECEIVED = date(2026, 9, 2)
DEADLINE = date(2026, 11, 3)
MEETING = '2026年11月16日10时'
FILES = {
    'COURT-ZHENXIN-GUIDE.pdf': '管理人债权申报指引',
    'COURT-ZHENXIN-CLAIM.docx': '管理人空白债权申报书',
    'COURT-ZHENXIN-REGISTER.docx': '管理人空白债权申报登记表',
    'COURT-ZHENXIN-LIST.docx': '管理人空白债权申报材料清单',
}


def register_font():
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    return 'STSong-Light'


def get_source_facts(input_dir):
    guide = PdfReader(input_dir / 'COURT-ZHENXIN-GUIDE.pdf')
    guide_text = '\n'.join(page.extract_text() or '' for page in guide.pages)
    claim = Document(input_dir / 'COURT-ZHENXIN-CLAIM.docx')
    claim_text = '\n'.join(p.text for p in claim.paragraphs)
    reg = Document(input_dir / 'COURT-ZHENXIN-REGISTER.docx')
    reg_text = '\n'.join(p.text for p in reg.paragraphs)
    list_doc = Document(input_dir / 'COURT-ZHENXIN-LIST.docx')
    list_text = '\n'.join(p.text for p in list_doc.paragraphs)
    all_text = guide_text + claim_text + reg_text + list_text
    compact = ''.join(all_text.split())
    required = ['2026年9月2日', '2026年11月3日前', '2026年9月2日（不含当天）',
                '2026年11月16日10时', '每年按360天、每月按30天计算']
    missing = [value for value in required if ''.join(value.split()) not in compact]
    if missing:
        raise ValueError('Trusted originals missing expected fact anchors: ' + repr(missing))
    return guide_text, claim_text, reg_text, list_text


def style_sheet(font):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='CNTitle', parent=styles['Title'], fontName=font,
                              fontSize=17, leading=23, alignment=TA_CENTER, spaceAfter=8*mm))
    styles.add(ParagraphStyle(name='CNBody', parent=styles['BodyText'], fontName=font,
                              fontSize=9.5, leading=15, spaceAfter=3*mm, wordWrap='CJK'))
    styles.add(ParagraphStyle(name='CNHead', parent=styles['Heading2'], fontName=font,
                              fontSize=11, leading=16, spaceBefore=4*mm, spaceAfter=2*mm))
    styles.add(ParagraphStyle(name='CNSmall', parent=styles['BodyText'], fontName=font,
                              fontSize=8, leading=12, wordWrap='CJK'))
    return styles


def para(text, style):
    return Paragraph(text.replace('&', '&amp;').replace('\n', '<br/>'), style)


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('STSong-Light', 8)
    canvas.setFillColor(colors.grey)
    canvas.drawString(20*mm, 12*mm, '东莞市震鑫磁业有限公司破产清算案 | 工作草稿')
    canvas.drawRightString(190*mm, 12*mm, f'{doc.page}')
    canvas.restoreState()


def pdf_doc(path, title, story, styles, font):
    frame = Frame(20*mm, 19*mm, 170*mm, 260*mm, id='normal')
    doc = BaseDocTemplate(str(path), pagesize=A4, title=title,
                          leftMargin=20*mm, rightMargin=20*mm, topMargin=17*mm, bottomMargin=19*mm)
    doc.addPageTemplates([PageTemplate(id='main', frames=[frame], onPage=footer)])
    doc.build(story)


def build_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = '事项与口径'
    rows = [
        ['事项', '当前结论/工作口径', '来源及定位', '状态'],
        ['债务人', '东莞市震鑫磁业有限公司', '四份管理人原件标题', '已核原件'],
        ['破产受理日', RECEIVED, 'COURT-ZHENXIN-GUIDE.pdf，第1页', '已核原件'],
        ['计息截止日（不含）', RECEIVED, 'COURT-ZHENXIN-GUIDE.pdf，第4页；REGISTER备注', '已核原件'],
        ['申报截止日', DEADLINE, 'COURT-ZHENXIN-GUIDE.pdf，第1页', '已核原件'],
        ['首次债权人会议', MEETING, 'COURT-ZHENXIN-GUIDE.pdf，第6页', '已核原件'],
        ['提交方式', '邮件发送完整盖章扫描件后，邮寄纸质材料；以管理人收到纸质材料的时间和内容为准；保留邮寄底单', 'COURT-ZHENXIN-GUIDE.pdf，第1页', '已核原件'],
        ['纸质份数/顺序', '一式一份；材料清单→登记表→申报书→主体及授权材料→送达方式确认书→债权事实证据', 'COURT-ZHENXIN-GUIDE.pdf，第2页', '已核原件'],
        ['利息日数', '按实际日历日差；受理日不计入。指引规定年利率/日利率折算每年360天、每月30天。', 'COURT-ZHENXIN-GUIDE.pdf，第4页', '合同/生效文书待核'],
        ['简式利息公式', '本金 × 年利率 × 实际天数 ÷ 360；须先核合同或生效法律文书的利率、期间、复利/违约金依据', '指引第4页；管理人要求计算明细见第2页', '待债权凭证'],
        ['债权人及金额', '当前没有委托方债权卷宗，身份、债权本金、利息、费用、性质及受偿均未核实；不得填0作为无债权结论', '当前材料覆盖范围', '待委托方补充'],
        ['法律效果提示', '管理人登记表与申报书不一致时以登记表为准；申报不等于管理人确认/最终确权', 'REGISTER声明；清单签收声明', '已核原件'],
    ]
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:D{ws.max_row}'
    widths = [25, 72, 48, 24]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='305B63')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.column == 2 and isinstance(cell.value, date):
                cell.number_format = 'yyyy-mm-dd'
    ws.sheet_view.showGridLines = False

    calc = wb.create_sheet('计息明细')
    headers = ['债权/凭证编号', '来源凭证/定位', '计息起日（含）', '计息截止日（不含）',
               '本金（元）', '年利率（小数）', '实际天数', '按360日折算利息（元）',
               '依据状态', '复核备注']
    calc.append(headers)
    calc.append(['待补原始凭证', '未收到合同、对账、付款或生效法律文书', None, RECEIVED,
                 None, None, '=IF(OR(C2="",D2=""),"",D2-C2)',
                 '=IF(OR(E2="",F2="",G2=""),"",ROUND(E2*F2*G2/360,2))',
                 '待补', '不得以空白视为零；核验利率、起息日、本金变化和既有清偿后再申报'])
    for r in range(3, 23):
        calc.cell(r, 7, f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})')
        calc.cell(r, 8, f'=IF(OR(E{r}="",F{r}="",G{r}=""),"",ROUND(E{r}*F{r}*G{r}/360,2))')
    calc.freeze_panes = 'A2'
    calc.auto_filter.ref = 'A1:J22'
    for i, width in enumerate([22, 42, 18, 20, 18, 19, 14, 25, 20, 58], 1):
        calc.column_dimensions[get_column_letter(i)].width = width
    for row in calc.iter_rows(min_row=2, max_row=22):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for row in calc.iter_rows(min_row=2, max_row=22):
        for col in (3, 4): row[col-1].number_format = 'yyyy-mm-dd'
        for col in (5, 8): row[col-1].number_format = '#,##0.00;[Red]-#,##0.00'
        row[5].number_format = '0.0000%'
    calc['A25'] = '公式口径'
    calc['B25'] = '按实际日期差计算天数；截止日不含。简式利息=本金×年利率×实际天数÷360。逐笔核实本金余额及每笔偿付，不能用汇总值掩盖本金变动。'
    calc.merge_cells('B25:J25')
    calc['B25'].alignment = Alignment(wrap_text=True, vertical='top')
    calc.row_dimensions[25].height = 38
    for sheet in (ws, calc):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='305B63')
        sheet.sheet_view.showGridLines = False
    calc['D2'].number_format = 'yyyy-mm-dd'

    gaps = wb.create_sheet('补件清单')
    gaps.append(['类别', '应收材料/核查事项', '用途/关联申报项目', '当前状态', '收到后复核点'])
    items = [
        ['主体资格', '营业执照复印件、法定代表人身份证明原件及身份证复印件；提交日最新企业登记信息', '债权人身份/登记表', '待委托方补充', '主体名称、统一代码、法定代表人、盖章及最新信息'],
        ['代理权限', '授权委托书原件；代理人身份证明；律师代理另附律所公函及律师执业证复印件', '登记表代理信息/申报签署', '待委托方补充', '权限范围、期限、签署人与盖章'],
        ['送达信息', '送达方式确认书及联系人/地址/电话/邮箱', '管理人联系及程序送达', '待委托方补充', '确认可送达且变更后及时书面告知'],
        ['债权形成', '合同/协议、订单、发票/交付/验收/付款凭证及对账资料', '本金及债权发生原因', '待委托方补充', '合同主体、履行、余额与会计记录勾稽'],
        ['偿付/抵销', '每笔收款、还款、抵销通知和账务凭证', '扣减本金、避免重复申报', '待委托方补充', '日期、金额、抵扣顺序及是否已由担保人清偿'],
        ['利息/其他费用', '合同计息条款、利率变更、催告及费用凭证', '孳息/违约金/其他金额', '待委托方补充', '利率依据、起止、日数、受理日截止；不主张无依据费用'],
        ['担保', '担保合同、登记证明、担保物权属/价值及实现情况', '债权性质及担保金额', '待委托方补充', '担保范围、登记与担保物；申报和确认状态分开'],
        ['诉讼/仲裁/执行', '起诉/受理、裁判及生效证明、仲裁文书、执行受理/裁定/到款凭证', '生效/在诉状态、费用及已执行受偿', '待委托方补充', '案号、状态、未决主张、执行到款扣减及禁重复'],
        ['官方程序文件', '管理人指引、申报书、登记表、材料清单', '提交格式和截止口径', '官方原件已提供', '原件版本；提交时按通知一式一份及顺序整理'],
    ]
    for item in items:
        gaps.append(item)
    gaps.freeze_panes = 'A2'
    gaps.auto_filter.ref = f'A1:E{gaps.max_row}'
    for i, width in enumerate([20, 62, 38, 22, 54], 1):
        gaps.column_dimensions[get_column_letter(i)].width = width
    for row in gaps.iter_rows(min_row=2):
        for cell in row: cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in gaps[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='305B63')
    gaps.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = 'auto'
    wb.save(path)


def build_declaration(path, styles):
    S = styles
    story = [para('东莞市震鑫磁业有限公司破产清算案', S['CNHead']),
             para('债权申报书（待补资料草稿，未签署/未提交）', S['CNTitle'])]
    fields = [
        ['申报人', '____________________________'], ['统一社会信用代码/身份证号', '____________________________'],
        ['住所', '________________________________________________________'], ['债务人', '东莞市震鑫磁业有限公司'],
        ['申报债权总额（人民币）', '待委托方债权凭证及计算底稿核实后填写'], ['其中：本金', '待核实；当前无债权本金材料'],
        ['孳息/违约金/滞纳金', '待核实；须按合同/生效文书及指引口径逐笔计算'], ['诉讼费用', '待核实；需费用凭证及程序文书'],
        ['其他', '待核实；需说明性质及依据'], ['债权性质', '待核实；依据担保、优先权及债权性质材料判断'],
        ['计息截止', '2026年9月2日（不含当天；管理人指引）'], ['申报截止', '2026年11月3日（管理人指引）'], ['事实和理由', '当前仅收到管理人公开指引及空白表单，未收到申报人主体、授权、合同、履行、付款、担保、抵销或诉讼执行资料。债权是否成立、金额、性质及受偿情况尚待核验，本草稿不主张特定债权金额。材料补齐后应说明债权形成原因、合同履行情况、逐笔清偿、担保及程序状态，并与登记表保持一致。'],
        ['申报人（签字/盖章）', '____________________________'], ['委托代理人签名', '____________________________'],
        ['日期', '______年____月____日'],
    ]
    data = [[para(a, S['CNSmall']), para(b, S['CNSmall'])] for a, b in fields]
    table = Table(data, colWidths=[48*mm, 122*mm], repeatRows=0)
    table.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'STSong-Light'),
        ('GRID', (0,0), (-1,-1), .4, colors.HexColor('#777777')),
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#EAF0EF')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6), ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.extend([table, Spacer(1, 5*mm), para('本文件仅为待补资料工作草稿。管理人收件或登记不代表对债权真实性、合法性、关联性或最终金额的确认。', S['CNSmall'])])
    pdf_doc(path, '债权申报书草稿', story, styles, 'STSong-Light')


def build_evidence_index(path, styles):
    S = styles
    story = [para('债权申报材料与证据目录', S['CNTitle']),
             para('东莞市震鑫磁业有限公司破产清算案 | 债权人：待补 | 版本：工作底稿', S['CNBody']),
             para('按管理人指引的整理顺序列示。当前清单不是实际提交/签收清单；未统计未取得文件的份数和页数。复印件应与原件一致并由债权人或受托人签名/盖章，管理人可要求核验原件；邮寄申报后按通知配合原件核对。', S['CNBody'])]
    data = [[para(x, S['CNSmall']) for x in ['序号', '文件类别/拟收材料', '当前状态与核验点', '关联申报项目']]]
    rows = [
        ['1', '债权申报材料清单', '待填写、签章；按实际材料顺序、数量、页数及原/复印件登记', '纸质材料总目录'],
        ['2', '债权申报登记表（管理人空白原件已提供）', '须补债权人/代理/金额/性质/担保/清偿信息并签章；与申报书一致', '金额汇总及债权性质'],
        ['3', '债权申报书（管理人空白原件已提供）', '须补形成原因、履行、担保及受偿事实；机构债权人加盖公章', '事实和理由、金额'],
        ['4', '主体资料及授权材料', '待补：营业执照、法定代表人材料、最新登记信息；委托时授权书及代理证明', '主体、签署与代理权限'],
        ['5', '送达方式确认书', '待补联系地址、邮编、电话、邮箱、联系人', '管理人程序送达'],
        ['6', '债权形成及履行证据', '待补合同/协议、订单、交付/验收、收付款及对账凭证；原件待核', '债权成立与本金余额'],
        ['7', '偿付、抵销及担保材料', '待补逐笔清偿/抵销凭证、担保合同、登记及担保物资料', '余额扣减与有无财产担保'],
        ['8', '利息及其他金额依据', '待补计息约定、利率变更和费用凭证；工作簿逐笔计算并核对受理日截止', '孳息/违约金/费用'],
        ['9', '诉讼、仲裁及执行材料（如有）', '待核是否存在；提交受理/裁判、生效、执行及实际到款材料', '程序状态、费用及避免重复受偿'],
        ['10', '管理人程序原件（已提供，不是债权事实证据）', '指引 PDF、申报书/登记表/材料清单 DOCX 各1份；工作目录所用原件，提交清单应按实际复印/打印件重新点数', '格式、期限、金额口径'],
    ]
    for row in rows:
        data.append([para(str(x), S['CNSmall']) for x in row])
    table = Table(data, colWidths=[11*mm, 55*mm, 73*mm, 31*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ('FONTNAME', (0,0), (-1,-1), 'STSong-Light'),
        ('GRID', (0,0), (-1,-1), .35, colors.HexColor('#777777')),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#DCE9E7')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4), ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
    ]))
    story.extend([table, Spacer(1, 4*mm),
                  para('程序提示：管理人通知的申报截止日为2026年11月3日；纸质材料一式一份邮寄，先发送盖章扫描件，申报内容及时间以管理人收到的纸质材料为准。寄送底单应留存。当前原件数量不等于最终拟提交材料数量。', S['CNSmall']),
                  Spacer(1, 3*mm), para('提交人签章：________________　　实际提交日：____________　　签收信息：待收件后填写', S['CNSmall'])])
    pdf_doc(path, '债权申报材料与证据目录', story, styles, 'STSong-Light')


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input-data', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    get_source_facts(args.input_data)
    font = register_font()
    styles = style_sheet(font)
    build_workbook(args.output / 'claim_calculation.xlsx')
    build_declaration(args.output / 'claim_declaration.pdf', styles)
    build_evidence_index(args.output / 'supporting_evidence_index.pdf', styles)


if __name__ == '__main__':
    main()
