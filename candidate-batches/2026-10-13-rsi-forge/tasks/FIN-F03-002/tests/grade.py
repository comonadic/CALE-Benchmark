import argparse
import json
import re
import subprocess
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def pdf_text(path):
    return subprocess.check_output(["pdftotext", "-layout", str(path), "-"], text=True, errors="replace")


def number(text, label):
    m = re.search(r"^\s*" + re.escape(label) + r"\s+([\d,]+\.\d{2})", text, re.M)
    if not m:
        raise ValueError(label)
    return float(m.group(1).replace(",", ""))


def section(text, title, end):
    a = text.index(title)
    b = text.index(end, a + len(title))
    return text[a:b]


def verify(output, input_dir):
    checks = {}
    xlsx = output / "working_capital_review.xlsx"
    pdf = output / "credit_investigation.pdf"
    checks["xlsx_exists"] = xlsx.is_file() and xlsx.stat().st_size > 1000
    checks["pdf_exists"] = pdf.is_file() and len(PdfReader(str(pdf)).pages) >= 2
    pdftext = re.sub(r"\s+", " ", subprocess.check_output(["pdftotext", str(pdf), "-"], text=True)) if pdf.exists() else ""
    wb = load_workbook(xlsx, data_only=False) if xlsx.exists() else None
    if wb:
        expected_sheets = {"来源与口径", "周转测算", "需求情景", "待核实事项"}
        checks["xlsx_exists"] = checks["xlsx_exists"] and expected_sheets.issubset(wb.sheetnames)
        combined = "\n".join(str(c.value or "") for ws in wb for row in ws.iter_rows() for c in row) + pdftext
        source_rows = list(wb["来源与口径"].iter_rows(min_row=2, values_only=True))
        def sourced_value(label, period):
            return [row[1] for row in source_rows if row[0] == label and period in str(row[2])]
    else:
        combined = pdftext
        sourced_value = lambda label, period: []

    checks["period_and_units"] = ("元" in combined and "2025" in combined and "2026" in combined)
    checks["scope_group_vs_parent"] = all(k in combined for k in ["合并", "母公司"]) and "不把集团负债等同申请人债务" in combined

    annual_text = pdf_text(input_dir / "CN-1225203635.pdf")
    interim_text = pdf_text(input_dir / "CN-1225495975.pdf")
    abs_annual = section(annual_text, "1、合并资产负债表", "2、母公司资产负债表")
    abs_h1 = section(interim_text, "1、合并资产负债表", "2、母公司资产负债表")
    pl_annual = section(annual_text, "3、合并利润表", "4、母公司利润表")
    cf_h1 = section(interim_text, "5、合并现金流量表", "6、母公司现金流量表")
    checks["source_revenue_2025"] = round(number(pl_annual, "一、营业总收入"), 2) == 128468230914.96 and "128,468,230,914.96" in combined
    checks["source_cost_2025"] = round(number(pl_annual, "其中：营业成本"), 2) == 115193343614.54 and "115,193,343,614.54" in combined
    checks["source_ar_2025"] = round(number(abs_annual, "应收账款"), 2) == 2978763231.09 and "2,978,763,231.09" in combined
    checks["source_inventory_2025"] = "11,371,777,008.15" in combined
    checks["source_short_debt_h1"] = "3,601,033,153.61" in combined
    checks["source_op_cash_h1"] = "3,923,561,511.73" in combined

    if wb:
        turn = wb["周转测算"]
        checks["turnover_inputs"] = all(turn.cell(r, 2).value and turn.cell(r, 3).value and turn.cell(r, 4).value and str(turn.cell(r, 5).value).startswith("=AVERAGE") for r in range(2, 7))
        checks["turnover_formulas"] = all(str(turn.cell(r, 6).value).startswith("=") and str(turn.cell(r, 7).value).startswith("=360/") for r in range(2, 7))
        mapping = " ".join(str(turn.cell(r, 8).value) for r in range(2, 7))
        checks["turnover_mapping"] = all(s in mapping for s in ["营业收入", "营业成本", "平均应收账款", "平均存货"])
        scen = wb["需求情景"]
        scen_text = " ".join(str(c.value or "") for row in scen.iter_rows() for c in row)
        checks["working_capital_formula"] = all(s in scen_text for s in ["=360/B11", "=B2*(1-B4)*(1+B5)/B12", "待核实"])
        checks["annualization"] = "!B5*2" in scen_text and "!B6*2" in scen_text and "机械年化" in scen_text
        checks["profit_margin"] = "!B4/'来源与口径'!B2" in scen_text and "营业利润率" in scen_text
        checks["growth_assumption"] = "预计销售收入增长率" in scen_text and "不是客户预测" in scen_text
        checks["financing_deductions_missing"] = "未提供" in scen_text and "待核实" in scen_text and wb["需求情景"]["B15"].value is None and wb["需求情景"]["B16"].value is None and wb["需求情景"]["B17"].value is None
        pending = " ".join(str(c.value or "") for row in wb["待核实事项"].iter_rows() for c in row)
        checks["no_fabricated_applicant_inputs"] = all(s in pending for s in ["申请主体", "申请金额", "授信", "未提供"])
        checks["cash_constraints"] = "受限资产" in combined and "境外" in combined and "可自由支配资金" in combined
    else:
        for k in ["turnover_inputs", "turnover_formulas", "turnover_mapping", "working_capital_formula", "annualization", "profit_margin", "growth_assumption", "financing_deductions_missing", "no_fabricated_applicant_inputs", "cash_constraints"]:
            checks[k] = False
    checks["no_hedge_limit_as_debt"] = "额度是授权上限而非实际持仓、贷款或现金占 用" in pdftext or "额度是授权上限而非实际持仓、贷款或现金占用" in pdftext
    checks["repayment_evidence"] = "经营现金流量净额为正" in pdftext and "申请人偿债覆盖" in pdftext
    checks["regulatory_example_boundary"] = "附件方法演示" in pdftext and "适用于所有客户" in pdftext and "机械" in pdftext
    checks["role_separation"] = all(s in pdftext for s in ["桌面分析", "现场调查", "独立风险评价", "授信审批"])
    caps = []
    for cap_id in ["scope_group_vs_parent", "period_and_units", "no_fabricated_applicant_inputs", "no_hedge_limit_as_debt"]:
        if not checks.get(cap_id, False):
            cap = {"scope_group_vs_parent": .25, "period_and_units": .25, "no_fabricated_applicant_inputs": .5, "no_hedge_limit_as_debt": .5}[cap_id]
            caps.append(cap)
    return checks, caps


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--input-data", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--logs", type=Path, required=True)
    args = p.parse_args()
    args.logs.mkdir(parents=True, exist_ok=True)
    contract = json.loads((Path(__file__).parent / "assertions.json").read_text(encoding="utf-8"))
    try:
        checks, caps = verify(args.output, args.input_data)
    except Exception as exc:
        report = {"schema_version": "1.0", "status": "verifier_failure", "error": f"{type(exc).__name__}: {exc}", "reward": 0.0, "assertions": []}
        (args.logs / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (args.logs / "reward.json").write_text(json.dumps({"reward": 0.0}, indent=2) + "\n", encoding="utf-8")
        return
    assertions = contract["assertions"]
    results = []
    for item in assertions:
        ok = bool(checks.get(item["id"], False))
        results.append({"id": item["id"], "kind": item["kind"], "unit": item["unit"], "status": "passed" if ok else "failed", "score": 1 if ok else 0, "message": item["description"]})
    passed = sum(r["score"] for r in results)
    raw = passed / contract["scale"]
    reward = min([raw] + caps) if caps else raw
    status = "graded" if all(checks.values()) else "graded_candidate_defect"
    report = {"schema_version": "1.0", "status": status, "passed": passed, "total": contract["scale"], "raw_reward": raw, "critical_caps": caps, "reward": reward, "assertions": results}
    (args.logs / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (args.logs / "reward.json").write_text(json.dumps({"reward": reward}, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
